"""H11: home builder workbook, 'Materials' sheet first, job x day schedule on the second sheet. Clean."""
from common import *
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
import datetime as dt

FID, SLUG = "H11", "homebuilder-materials-sheet-then-job-grid-clean"
README = ("Clean schedule (no double-bookings, no one assigned while off) on the second sheet of an xlsx "
          "whose first sheet is a Materials order list with job names, dates and people's names (must not "
          "ask which sheet): job x day grid with merged title, a roofing subcontractor-only job, Helper "
          "placeholder, TBD, em dash and 'done' cells, Shop and OFF rows, nickname 'Gordo'.")


def build():
    dates = days("2026-10-12", 5)
    Mo, Tu, We, Th, Fr = dates
    EM = "—"
    L3, L7, KR, HD, MO = ("Lot 3 Sorrel Ridge", "Lot 7 Sorrel Ridge", "Kim remodel", "Halvorsen deck",
                          "Mt. Olive Church roof")
    NR = "Northside Roofing LLC"
    jobs = [
        (L3, ["Dave, Curtis, Jamie", "Dave, Curtis, Jamie", "Dave, Curtis, Jamie", "Dave, Curtis", "Dave, Curtis, Jamie"]),
        (L7, ["Raymond, Eddie, Gordo", "Raymond, Eddie, Gordo", "Raymond, Eddie, Helper", "Raymond, Eddie, Gordo",
              "Raymond, Eddie"]),
        (KR, ["Tobias, Lucas", "Tobias, Lucas", "Tobias, Nell", "Tobias, Nell", "TBD"]),
        (HD, [EM, "Sam, Nell", "Sam, Gordo", "Sam", "Sam, Nell"]),
        (MO, [NR, NR, "", NR, "done"]),
        ("Shop", ["Nell", "", "Lucas", "Lucas", "Gordo"]),
    ]
    off_row = ["", "", "", "Jamie", "Tobias, Lucas"]
    recs, inactive = [], []
    for job, cells in jobs:
        for d, c in zip(dates, cells):
            if c and norm(c).lower() in INACTIVE:
                inactive.append([job, d])
            recs += names_cell(job, d, c)
    for d, c in zip(dates, off_row):
        recs += names_cell("OFF", d, c, status="off")

    wb = Workbook()
    mat = wb.active
    mat.title = "Materials"
    hdr = ["Item", "Qty", "For job", "Supplier", "Need by", "Ordered by"]
    for j, v in enumerate(hdr, start=1):
        mat.cell(1, j, v).font = Font(bold=True)
    items = [
        ("2x10x16 SPF #2", 48, L3, "Hillman Lumber", "2026-10-13", "Dave"),
        ("1/2\" CDX ply", 60, L7, "Hillman Lumber", "2026-10-14", "Raymond"),
        ("LUS210 hangers", "2 boxes", L3, "Hillman Lumber", "2026-10-13", "Curtis"),
        ("Base cabinets (12 pcs)", 1, KR, "Cabinet Depot", "2026-10-15", "Tobias"),
        ("Composite decking 16'", 40, HD, "Coastal Supply", "2026-10-13", "Sam"),
        ("1/2\" drywall 4x12", 30, KR, "Coastal Supply", "2026-10-16", "Nell"),
        ("Joist tape", 3, HD, "Coastal Supply", "2026-10-13", "Sam"),
        ("Ice & water shield", 6, MO, "Northside to supply", "", ""),
        ("Construction adhesive", 24, L7, "Hillman Lumber", "2026-10-14", "Eddie"),
    ]
    for i, row in enumerate(items, start=2):
        for j, v in enumerate(row, start=1):
            if j == 5 and v:
                c = mat.cell(i, j, D(v))
                c.number_format = "m/d/yyyy"
            elif v != "":
                mat.cell(i, j, v)
    mat.cell(len(items) + 3, 1, "Call orders in by 2pm for next-day delivery.")
    for col, w in zip("ABCDEF", [26, 10, 22, 20, 12, 12]):
        mat.column_dimensions[col].width = w

    ws = wb.create_sheet("Schedule")
    ws.cell(1, 1, "ALDER CREEK HOMES — CREW SCHEDULE").font = Font(bold=True, size=16)
    ws.merge_cells("A1:F1")
    ws.cell(2, 1, "Week of 10/12/2026").font = Font(bold=True)
    ws.merge_cells("A2:F2")
    hdr = ["Job / Address"] + [f"{n} {md(d)}" for n, d in zip(["Mon", "Tue", "Wed", "Thu", "Fri"], dates)]
    for j, v in enumerate(hdr, start=1):
        c = ws.cell(4, j, v)
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="DDEBF7")
    r = 5
    for job, cells in jobs:
        ws.cell(r, 1, job)
        for j, c in enumerate(cells):
            if c:
                ws.cell(r, 2 + j, c)
        r += 1
    ws.cell(r, 1, "OFF").font = Font(bold=True)
    for j, c in enumerate(off_row):
        if c:
            ws.cell(r, 2 + j, c)
    ws.cell(r + 2, 1, "Lot 3 walk-through with the buyers Thu 4pm.")
    ws.column_dimensions["A"].width = 24
    for col in "BCDEF":
        ws.column_dimensions[col].width = 24
    d = fixture_dir(FID, SLUG)
    wb.save(os.path.join(d, "input.xlsx"))

    flags = [
        flag_nb("Sam", Mo),
        flag_un(KR, Fr),
        flag_un(MO, We),
    ]
    expected = dict(workdays=MON_FRI, layout="job_by_day", assignments=recs, flags=flags, info=[],
                    must_not_ask=[{"type": "pick_sheet"}, {"type": "weekday_mismatch"}])
    meta = dict(window=dates, roster=None, person_canon={}, internal=["Shop"], placeholders=["Helper"], subs=[NR],
                jobs=[j for j, _ in jobs], inactive=inactive, pseudo_rows=["OFF"])
    return README, "input.xlsx", expected, meta
