"""H02: builder's person x day grid in xlsx with crew sections, a floater listed in two sections."""
from common import *
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

FID, SLUG = "H02", "builder-person-grid-crew-sections"
README = ("Custom-home builder's xlsx person x day grid: merged title block, 'Week of' line, 'Mon 10/19' "
          "headers, a blank spacer column, merged crew-section label rows (FRAMING / FINISH), a floater "
          "listed under two sections, two sites in one cell, shop plus a site, notes in cells, PTO/sick/"
          "OFF/vac, a Saturday column under Mon-Fri, a placeholder Helper row and a totals row.")


def build():
    dates = days("2026-10-19", 6)  # Mon..Sat
    Mo, Tu, We, Th, Fr, Sa = dates
    P, L, C, H, A = "Pruitt Res.", "Lot 14 Deer Run", "Cole St duplex", "Hendricks kitchen", "Ambrose bath"
    sections = [
        ("FRAMING", [
            ("Rich Kowalczyk", [P, P, P, P, P, P]),
            ("Danny Ochoa", [P, P, P, P, P, P]),
            ("Kev Tran", [P, P, P, L, L, ""]),
            ("Shane Mulroney", [C, C, L, L, L, ""]),
            ("Eli Bowman", ["PTO", L, L, L, L, ""]),
            ("Helper", [P, P, L, L, "", ""]),
        ]),
        ("FINISH / REMODEL", [
            ("Gary Linde", ["Shop / " + H, H, H + " - pick up cabinets", H, H, ""]),
            ("Pat Oyelaran", [A, A, "sick", A + " (punch list)", "OFF", ""]),
            ("Marisol Vega", [H, H, H, "vac", "vac", ""]),
            ("Tom Brandt", ["Shop", A, H, H + " / " + A, H, ""]),
            ("Kev Tran", ["", "", H, "", "", ""]),
            ("Jake Furlong", ["Shop - orientation", "", "", P, P, ""]),
        ]),
    ]

    recs = []
    for _, people in sections:
        for name, cells in people:
            for d, c in zip(dates, cells):
                recs += person_cell(name, d, c)

    totals = []
    for d in dates:
        totals.append(sum(1 for r in recs if r["date"] == d and "site" in r))

    wb = Workbook()
    ws = wb.active
    ws.title = "Schedule"
    ncol = 8  # A name, B spacer, C..H Mon..Sat
    ws.cell(1, 1, "HOLLIS & GRANT BUILDERS — CREW SCHEDULE").font = Font(bold=True, size=16)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncol)
    ws.cell(2, 1, "Week of 10/19/2026").font = Font(italic=True, size=12)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncol)
    ws.cell(1, 1).alignment = Alignment(horizontal="center")
    ws.cell(2, 1).alignment = Alignment(horizontal="center")
    hdr = ["Name", None] + [f"{n} {md(d)}" for n, d in zip(["Mon", "Tue", "Wed", "Thu", "Fri", "Sat"], dates)]
    for j, v in enumerate(hdr, start=1):
        if v is not None:
            c = ws.cell(4, j, v)
            c.font = Font(bold=True)
            c.fill = PatternFill("solid", fgColor="D9D9D9")
    r = 5
    for label, people in sections:
        ws.cell(r, 1, label).font = Font(bold=True, color="FFFFFF")
        ws.cell(r, 1).fill = PatternFill("solid", fgColor="44546A")
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ncol)
        r += 1
        for name, cells in people:
            ws.cell(r, 1, name)
            for j, c in enumerate(cells):
                if c:
                    ws.cell(r, 3 + j, c)
            r += 1
        r += 1  # blank spacer row between sections
    ws.cell(r, 1, "Total on jobs").font = Font(bold=True)
    for j, t in enumerate(totals):
        ws.cell(r, 3 + j, t).font = Font(bold=True)
    r += 2
    ws.cell(r, 1, "Framing inspection at Pruitt Thu 10am - Rich meets the inspector.")
    ws.cell(r + 1, 1, "Cabinets for Hendricks arrive Wed. Call the office with any changes.")
    ws.cell(r + 3, 1, "Posted by: ____________________")
    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 3
    for col in "CDEFGH":
        ws.column_dimensions[col].width = 26
    d = fixture_dir(FID, SLUG)
    wb.save(os.path.join(d, "input.xlsx"))

    flags = [
        flag_db("Kev Tran", We, [P, H]),
        flag_db("Tom Brandt", Th, [H, A]),
        flag_nb("Jake Furlong", Tu),
        flag_nb("Jake Furlong", We),
        flag_un(A, We),
    ]
    infos = [info("internal_plus_site", "Gary Linde", Mo, ["Shop", H])]
    expected = dict(workdays=MON_FRI, layout="person_by_day", assignments=recs, flags=flags, info=infos,
                    must_not_ask=[{"type": "weekday_mismatch"}, {"type": "fill_down"}])
    meta = dict(window=dates, roster=None, person_canon={}, internal=["Shop"], placeholders=["Helper"], subs=[])
    return README, "input.xlsx", expected, meta
