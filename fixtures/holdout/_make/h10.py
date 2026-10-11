"""H10: masonry company's October job calendar (calendar-month grid) in xlsx."""
from common import *
import datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

FID, SLUG = "H10", "mason-october-calendar-month"
README = ("Masonry company's October 2026 job calendar in xlsx (Excel calendar template style): Sunday-first "
          "week rows, day-number row over 'Job: names' entry lines, blank days from September, OFF/PTO/VAC "
          "lines, Saturday overtime with a note, Yard as internal, crews changing jobs mid-month, merged "
          "title and footer notes.")


def build():
    G, B, A, P, Y = ("Greer retaining wall", "Bell Tower Lofts", "Ashford chimney", "Pinehurst Library steps",
                     "Yard")
    e = {}

    def put(day, *lines):
        e[f"2026-10-{day:02d}"] = list(lines)

    put(1, f"{A}: Walt, Deon", f"{B}: Sal, Rusty, Andy", f"{Y}: Hai, Victor")
    put(2, f"{A}: Walt, Deon", f"{B}: Sal, Rusty, Andy", f"{Y}: Hai")
    put(3, f"{B}: Sal (OT)")
    put(5, f"{A}: Walt, Deon, Hai", f"{B}: Sal, Rusty, Andy, Victor")
    put(6, f"{A}: Walt, Deon, Hai", f"{B}: Sal, Rusty, Andy, Victor")
    put(7, f"{A}: Walt, Deon", f"{B}: Sal, Rusty, Andy, Victor", f"{Y}: Hai")
    put(8, f"{A}: Walt, Deon, Hai", f"{B}: Sal, Rusty, Victor", "OFF: Andy")
    put(9, f"{A}: Walt, Deon, Hai", f"{B}: Sal, Rusty, Andy, Victor")
    put(12, f"{G}: Walt, Deon, Hai", f"{B}: Sal, Rusty, Andy, Victor")
    put(13, f"{G}: Walt, Deon, Hai", f"{B}: Sal, Rusty, Andy, Victor")
    put(14, f"{G}: Walt, Deon, Hai", f"{B}: Rusty, Andy, Victor", "PTO: Sal")
    put(15, f"{G}: Walt, Deon, Hai", f"{B}: Sal, Rusty, Andy, Victor")
    put(16, f"{G}: Walt, Deon, Hai, Rusty", f"{B}: Sal, Rusty, Andy, Victor")
    put(19, f"{G}: Walt, Deon, Hai", f"{B}: Sal, Rusty", f"{P}: Andy, Victor")
    put(20, f"{G}: Walt, Deon, Hai", f"{B}: Sal, Rusty", f"{P}: Andy, Victor")
    put(21, f"{G}: Walt, Deon, Hai", f"{P}: Andy, Victor, Sal, Rusty")
    put(22, f"{G}: Walt, Deon, Hai", f"{B}: Sal, Rusty", f"{P}: Andy, Victor")
    put(23, f"{G}: Walt, Deon", f"{B}: Sal, Rusty", f"{P}: Andy, Victor", f"{Y}: Hai")
    put(26, f"{G}: Walt, Deon, Hai", f"{P}: Andy, Victor, Sal, Rusty")
    put(27, f"{G}: Walt, Deon, Hai", f"{P}: Andy, Victor, Sal, Rusty")
    put(28, f"{G}: Walt, Hai", f"{P}: Andy, Victor, Sal, Rusty", "VAC: Deon")
    put(29, f"{G}: Walt, Hai, Deon", f"{P}: Andy, Victor, Sal", "VAC: Deon")
    put(30, f"{G}: Walt, Hai", f"{P}: Andy, Victor, Sal, Rusty", "VAC: Deon")

    recs = []
    for d, lines in e.items():
        for line in lines:
            label, names = line.split(":", 1)
            st = STATUS.get(label.strip().lower())
            if st:
                recs += names_cell(label, d, names, status=st)
            else:
                recs += names_cell(label.strip(), d, names)

    wb = Workbook()
    ws = wb.active
    ws.title = "Oct 2026"
    ws.cell(1, 1, "COBB & SONS MASONRY — JOB CALENDAR").font = Font(bold=True, size=16)
    ws.merge_cells("A1:G1")
    ws.cell(2, 1, "OCTOBER 2026").font = Font(bold=True, size=14)
    ws.merge_cells("A2:G2")
    for c in (ws.cell(1, 1), ws.cell(2, 1)):
        c.alignment = Alignment(horizontal="center")
    for j, n in enumerate(["Sun", "Mon", "Tue", "Wed", "Thu", "Fri", "Sat"], start=1):
        c = ws.cell(4, j, n)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="7F6000")
        c.alignment = Alignment(horizontal="center")
    start = D("2026-09-27")
    r = 5
    for wk in range(5):
        for j in range(7):
            day = start + dt.timedelta(wk * 7 + j)
            if day.month == 10:
                c = ws.cell(r, 1 + j, day.day)
                c.font = Font(bold=True)
                c.fill = PatternFill("solid", fgColor="FFF2CC")
                for i, line in enumerate(e.get(day.isoformat(), [])):
                    ws.cell(r + 1 + i, 1 + j, line).alignment = Alignment(wrap_text=True, vertical="top")
        r += 5
    r += 1
    ws.cell(r, 1, "Walt has the final say on rain days. Mortar delivery to Greer on Mondays.")
    ws.cell(r + 1, 1, "Bell Tower punch list wraps up 10/23 - Pinehurst after that.")
    for col in "ABCDEFG":
        ws.column_dimensions[col].width = 30
    d = fixture_dir(FID, SLUG)
    wb.save(os.path.join(d, "input.xlsx"))

    def dd(n):
        return f"2026-10-{n:02d}"

    flags = [
        flag_db("Rusty", dd(16), [G, B]),
        flag_awo("Deon", dd(29), [G], "vacation"),
        flag_nb("Victor", dd(2)),
        flag_nb("Rusty", dd(29)),
        flag_un(B, dd(21)),
    ]
    infos = []
    expected = dict(workdays=MON_FRI, layout="calendar_month", assignments=recs, flags=flags, info=infos)
    window = days("2026-10-01", 30)
    meta = dict(window=window, roster=None, person_canon={}, internal=[Y], placeholders=[], subs=[])
    return README, "input.xlsx", expected, meta
