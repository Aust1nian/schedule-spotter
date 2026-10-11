"""H12: exteriors company, person x day, two weeks side by side in xlsx, Mon-Sat."""
from common import *
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

FID, SLUG = "H12", "exteriors-two-weeks-side-by-side"
README = ("Siding/windows company's xlsx person x day board with two weeks side by side (Mon-Sat), merged "
          "'WEEK OF' bands over single-letter day headers (M T W Th F Sa + M/D), a blank spacer column "
          "between weeks, two different Mikes by full name, a two-site cell, Shop plus a site, off/x/PTO/"
          "vac/sick, a Saturday left blank and a job left empty mid-week.")


def build():
    w1 = days("2026-10-12", 6)
    w2 = days("2026-10-19", 6)
    F, MG, HW, SP, OS = ("Fenton Rd siding", "Maple Grove Condos", "Haley windows", "St. Paul's gutters",
                         "Ostrander re-side")
    rows = [
        ("Travis Lunde", [F] * 6, [F] * 5 + ["off"]),
        ("Mike Doyle", [F] * 5 + [""], [F, F, F, MG, F, "off"]),
        ("Mike Szabo", [MG] * 6, [MG] * 6),
        ("Kyle Brenner", [MG] * 4 + ["PTO", "PTO"], [MG] * 5 + ["off"]),
        ("Ray Castillo", [MG] * 6, [MG, MG, MG, "PTO", MG, "off"]),
        ("Juan Pablo Rios", [HW, HW, HW, f"{HW} / {SP}", SP, "off"], [SP, SP, HW, SP, HW, "off"]),
        ("Dmitri Volkov", [HW] * 5 + ["off"], [HW] * 6),
        ("Ollie Grant", [SP] * 5 + ["x"], [SP, SP, "Shop", "Shop", "", "off"]),
        ("Pedro Aguilar", [OS] * 6, [OS] * 5 + ["off"]),
        ("Nick Hamm", [OS] * 5 + ["off"], [OS] * 5 + ["x"]),
        ("Jesse Ward", ["Shop"] + [OS] * 4 + ["off"], [OS] * 5 + ["off"]),
        ("Tommy Vu", ["vac"] * 6, [f"Shop / {F}", F, F, F, F, "off"]),
        ("Bobby Kerr", [F] * 6, [F] * 5 + ["off"]),
        ("Chris Delaney", [HW, HW, "sick", HW, HW, "x"], [HW] * 5 + ["off"]),
    ]
    recs = []
    for name, c1, c2 in rows:
        for d, c in zip(w1 + w2, c1 + c2):
            recs += person_cell(name, d, c)

    wb = Workbook()
    ws = wb.active
    ws.title = "Board"
    ws.cell(1, 1, "SUMMIT RIDGE EXTERIORS").font = Font(bold=True, size=16)
    ws.merge_cells("A1:N1")
    band = PatternFill("solid", fgColor="F8CBAD")
    ws.cell(3, 2, "WEEK OF 10/12/2026").font = Font(bold=True)
    ws.cell(3, 2).fill = band
    ws.cell(3, 2).alignment = Alignment(horizontal="center")
    ws.merge_cells("B3:G3")
    ws.cell(3, 9, "WEEK OF 10/19/2026").font = Font(bold=True)
    ws.cell(3, 9).fill = band
    ws.cell(3, 9).alignment = Alignment(horizontal="center")
    ws.merge_cells("I3:N3")
    letters = ["M", "T", "W", "Th", "F", "Sa"]
    ws.cell(4, 1, "Name").font = Font(bold=True)
    for j, (l, d) in enumerate(zip(letters, w1)):
        ws.cell(4, 2 + j, f"{l} {md(d)}").font = Font(bold=True)
    for j, (l, d) in enumerate(zip(letters, w2)):
        ws.cell(4, 9 + j, f"{l} {md(d)}").font = Font(bold=True)
    r = 5
    for name, c1, c2 in rows:
        ws.cell(r, 1, name)
        for j, c in enumerate(c1):
            if c:
                ws.cell(r, 2 + j, c)
        for j, c in enumerate(c2):
            if c:
                ws.cell(r, 9 + j, c)
        r += 1
    r += 1
    ws.cell(r, 1, "Ostrander dumpster swap Tue 10/20. Haley window order ships Thu 10/15.")
    ws.column_dimensions["A"].width = 18
    for col in "BCDEFGIJKLMN":
        ws.column_dimensions[col].width = 18
    ws.column_dimensions["H"].width = 3
    d = fixture_dir(FID, SLUG)
    wb.save(os.path.join(d, "input.xlsx"))

    flags = [
        flag_db("Juan Pablo Rios", w1[3], [HW, SP]),
        flag_nb("Mike Doyle", w1[5]),
        flag_nb("Ollie Grant", w2[4]),
        flag_un(SP, w2[2]),
    ]
    infos = [info("internal_plus_site", "Tommy Vu", w2[0], ["Shop", F])]
    expected = dict(workdays=MON_SAT, layout="person_by_day", arrangement="side_by_side", assignments=recs,
                    flags=flags, info=infos,
                    must_not_ask=[{"type": "same_person", "a": "Mike Doyle", "b": "Mike Szabo"},
                                  {"type": "weekday_mismatch"}])
    meta = dict(window=w1 + w2, roster=None, person_canon={}, internal=["Shop"], placeholders=[], subs=[])
    return README, "input.xlsx", expected, meta
