"""H14: restoration company's week typed as a table in Word (.docx), person x day. Clean."""
from common import *
import docx
from docx.shared import Pt

FID, SLUG = "H14", "restoration-word-table-clean"
README = ("Clean schedule (no double-bookings, no one assigned while off) typed as a table in a Word .docx: "
          "heading and 'week of 10/26/2026' paragraph above, full weekday names (Monday..Friday) as the only "
          "column headers, Office as an internal site and Office plus a job, notes after ' - ' in cells, "
          "PTO/vac, blank days, a job with nobody on it mid-week and a signature line below.")


def build():
    dates = days("2026-10-26", 5)
    Mo, Tu, We, Th, Fr = dates
    DA, PC, LC = "Delmar Ave water loss", "Pike Creek fire job", "Unit 12 Laurel Ct mold"
    rows = [
        ("Jared Quinn", [DA, DA, DA, PC, PC]),
        ("Lena Fox", [DA, DA, "Office", f"Office / {DA}", DA]),
        ("Tomas Herrera", [PC, PC, "Office - training", PC, "PTO"]),
        ("Bryce Neal", [LC, LC, f"{LC} - containment", "", LC]),
        ("Ade Okafor", [LC, LC, "", LC, LC]),
        ("Rocco Bianchi", [PC, PC, DA, DA, DA]),
        ("Shawn Pratt", ["vac"] * 5),
    ]
    recs = []
    for name, cells in rows:
        for d, c in zip(dates, cells):
            recs += person_cell(name, d, c)

    doc = docx.Document()
    doc.add_heading("Hearthstone Restoration", level=1)
    p = doc.add_paragraph("Field schedule — week of 10/26/2026")
    p.runs[0].font.size = Pt(13)
    doc.add_paragraph("Moisture readings go in the job folder every day. Trucks back at the office by 4.")
    table = doc.add_table(rows=1 + len(rows), cols=6)
    table.style = "Table Grid"
    for j, h in enumerate(["Tech", "Monday", "Tuesday", "Wednesday", "Thursday", "Friday"]):
        table.cell(0, j).text = h
        table.cell(0, j).paragraphs[0].runs[0].bold = True
    for i, (name, cells) in enumerate(rows, start=1):
        table.cell(i, 0).text = name
        for j, c in enumerate(cells, start=1):
            table.cell(i, j).text = c
    doc.add_paragraph("")
    doc.add_paragraph("Questions? Call Marla at the office.")
    doc.add_paragraph("Signed: ______________________   (ops manager)")
    d = fixture_dir(FID, SLUG)
    doc.save(os.path.join(d, "input.docx"))

    flags = [
        flag_nb("Bryce Neal", Th),
        flag_nb("Ade Okafor", We),
        flag_un(PC, We),
    ]
    infos = [info("internal_plus_site", "Lena Fox", Th, ["Office", DA])]
    expected = dict(workdays=MON_FRI, layout="person_by_day", assignments=recs, flags=flags, info=infos)
    meta = dict(window=dates, roster=None, person_canon={}, internal=["Office"], placeholders=[], subs=[])
    return README, "input.docx", expected, meta
