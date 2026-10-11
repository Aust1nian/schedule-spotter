"""H03: remodeler's job x day grid in xlsx, several names per cell, near-identical addresses."""
from common import *
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

FID, SLUG = "H03", "remodeler-job-grid-twin-addresses"
README = ("Remodeler's xlsx job x day grid: merged title, names split by comma/&/'/'/+/line breaks, "
          "1420 vs 1240 Johnson St as two real jobs (ask, never merge), Mike B vs Mike R, an electrical "
          "subcontractor at two jobs, Helper placeholder, TBD/open/done cells, a blank job-day, Shop row, "
          "OFF and VAC pseudo-rows (one with a note), footer notes.")


def build():
    dates = days("2026-10-26", 5)
    Mo, Tu, We, Th, Fr = dates
    jobs = [
        ("1420 Johnson St", ["Brian, Luke", "Brian, Luke", "Brian / Luke / Helper", "Brian, Luke", "Brian"]),
        ("1240 Johnson St", ["Manny & Rosa", "Manny & Rosa", "Manny", "TBD", "Manny & Rosa"]),
        ("Whitfield addition", ["Kenji\nTrevor\nIke", "Kenji\nTrevor\nIke", "Kenji\nTrevor",
                                "Kenji\nTrevor\nIke\nMike R", "Kenji\nTrevor"]),
        ("88 Larchmont Ave", ["Dom, Nate", "Dom, Nate, Delta Electric LLC", "", "Dom, Nate", "open"]),
        ("Fenwick basement", ["Phil + Mike R", "Phil + Mike R", "Phil, Rosa, Mike R", "Phil + Mike R", "done"]),
        ("Riverside Dental TI", ["Mike B", "Mike B, Delta Electric LLC", "Mike B + Ike", "Delta Electric LLC",
                                 "Mike B, Nate"]),
        ("Shop", ["", "", "Dom", "Luke", "Ike, Trevor, Mike R"]),
    ]
    pseudo = [
        ("OFF", "off", ["", "Kenji (dentist)", "", "", "Luke"]),
        ("VAC", "vacation", ["", "", "", "Rosa", ""]),
    ]
    recs = []
    inactive = []
    for job, cells in jobs:
        for d, c in zip(dates, cells):
            if c and norm(c).lower() in INACTIVE:
                inactive.append([job, d])
            recs += names_cell(job, d, c)
    for label, st, cells in pseudo:
        for d, c in zip(dates, cells):
            recs += names_cell(label, d, c, status=st)

    wb = Workbook()
    ws = wb.active
    ws.title = "Week"
    ws.cell(1, 1, "BRIGHTWATER REMODELING").font = Font(bold=True, size=16)
    ws.merge_cells("A1:F1")
    ws.cell(2, 1, "Schedule — week of Oct 26, 2026").font = Font(size=12)
    ws.merge_cells("A2:F2")
    hdr = ["Job"] + [f"{n} {md(d)}" for n, d in zip(["Mon", "Tue", "Wed", "Thu", "Fri"], dates)]
    for j, v in enumerate(hdr, start=1):
        c = ws.cell(4, j, v)
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="BDD7EE")
    r = 5
    for job, cells in jobs:
        ws.cell(r, 1, job).font = Font(bold=True)
        for j, c in enumerate(cells):
            if c:
                cell = ws.cell(r, 2 + j, c)
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        r += 1
    for label, _, cells in pseudo:
        ws.cell(r, 1, label).font = Font(bold=True, color="C00000")
        for j, c in enumerate(cells):
            if c:
                ws.cell(r, 2 + j, c)
        r += 1
    r += 1
    ws.cell(r, 1, "Delta Electric LLC - rough-in inspection at Riverside Thu, someone has to let them in.")
    ws.cell(r + 1, 1, "Rosa's vacation day is Thu 10/29. Dumpster at Fenwick gets pulled Fri.")
    ws.column_dimensions["A"].width = 24
    for col in "BCDEF":
        ws.column_dimensions[col].width = 22
    d = fixture_dir(FID, SLUG)
    wb.save(os.path.join(d, "input.xlsx"))

    flags = [
        flag_db("Mike R", Th, ["Whitfield addition", "Fenwick basement"]),
        flag_awo("Kenji", Tu, ["Whitfield addition"], "off"),
        flag_nb("Manny", Th),
        flag_nb("Dom", Fr),
        flag_nb("Nate", We),
        flag_nb("Phil", Fr),
        flag_nb("Mike B", Th),
        flag_un("1240 Johnson St", Th),
        flag_un("88 Larchmont Ave", Fr),
        flag_un("88 Larchmont Ave", We),
    ]
    infos = [
        info("internal_plus_site", "Luke", Th, ["Shop", "1420 Johnson St"]),
        info("internal_plus_site", "Trevor", Fr, ["Shop", "Whitfield addition"]),
        info("subcontractor_multi_site", "Delta Electric LLC", Tu, ["88 Larchmont Ave", "Riverside Dental TI"]),
    ]
    q_site = {"type": "same_site", "a": "1420 Johnson St", "b": "1240 Johnson St"}
    expected = dict(
        workdays=MON_FRI, layout="job_by_day", assignments=recs, flags=flags, info=infos,
        ask=[q_site],
        must_not_ask=[{"type": "same_person", "a": "Mike B", "b": "Mike R"}, {"type": "weekday_mismatch"}],
        answers=[{"question": q_site, "answer": "no"}],
        after_answers={"flags": flags},
    )
    meta = dict(window=dates, roster=None, person_canon={}, internal=["Shop"], placeholders=["Helper"],
                subs=["Delta Electric LLC"], jobs=[j for j, _ in jobs], inactive=inactive, pseudo_rows=["OFF", "VAC"])
    return README, "input.xlsx", expected, meta
