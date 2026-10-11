"""Category C: layouts.
Run: uv run --with openpyxl --with odfpy --with python-docx python3 build_C.py"""
import calendar
import datetime as dt

from fxlib import *

WEEK = days("10-12", 5)
WEEK2 = days("10-19", 5)
J, S, O, P = "1420 Johnson St", "Smith remodel", "Oak Ridge Apts", "Parker kitchen"
NO_LAYOUT_QS = [Q("pick_layout"), Q("column_role")]
FIXED = dt.datetime(2026, 10, 9, 12, 0)


def new_workbook():
    import openpyxl
    wb = openpyxl.Workbook()
    wb.properties.creator = "Office"
    wb.properties.created = wb.properties.modified = FIXED
    return wb


def write_person_sheet(ws, rows, dates, header="Name"):
    ws.cell(row=1, column=1, value=header)
    for i, x in enumerate(dates):
        c = ws.cell(row=1, column=2 + i, value=pydate(x))
        c.number_format = "m/d/yyyy"
    for r, (person, cells) in enumerate(rows, start=2):
        ws.cell(row=r, column=1, value=person)
        for i, cell in enumerate(cells):
            if cell:
                ws.cell(row=r, column=2 + i, value=cell)


def c01():
    rows = [
        ("Mike Baker", [J] * 5),
        ("Jose Garcia", [S] * 5),
        ("Luis Ortiz", [O, O, S, O, O]),
        ("Dave Kim", [J] * 5),
        ("Ana Ruiz", [O, O, S, O, O]),
    ]
    lines = [["Date", "Name", "Job"]] + grid_to_list_rows(rows, WEEK, mdyyyy)
    exp = expected("C", "input.csv", layout="long_list", arrangement="single",
                   assignments=person_grid(rows, WEEK),
                   flags=[UF(O, "10-14")], info=[], must_not_ask=NO_LAYOUT_QS)
    emit("C01", "long-list", "Long list, one row per assignment (Date, Name, Job); Oak Ridge has a mid-week gap, so it is unstaffed Wednesday.",
         {"input.csv": csv_bytes(lines)}, exp)


def c02():
    rows = [
        ("Mike Baker", [J, J, J, f"{J} / {P}", J]),
        ("Jose Garcia", [S] * 5),
        ("Luis Ortiz", [O] * 5),
        ("Dave Kim", [P] * 5),
        ("Tom Reed", [S, S, "OFF", S, S]),
        ("Ana Ruiz", [O, O, O, P, P]),
    ]
    lines = [["Name"] + [mdyyyy(x) for x in WEEK]] + [[p] + c for p, c in rows]
    exp = expected("C", "input.csv", layout="person_by_day", arrangement="single",
                   assignments=person_grid(rows, WEEK),
                   flags=[DB("Mike Baker", "10-15", J, P)], info=[], must_not_ask=NO_LAYOUT_QS)
    emit("C02", "person-by-day", "Person x day grid: people down the side, dates across, site names in cells.",
         {"input.csv": csv_bytes(lines)}, exp, {"window": WEEK})


def c03():
    rows = [
        (J, ["Mike & Jose", "Mike & Jose", "Mike, Jose, Tom", "Mike and Tom", "Mike / Tom"]),
        (S, ["Luis + Dave", "Luis, Dave", "Luis", "", "Luis & Ana"]),
        (O, ["Tom, Ana", "Tom & Ana", "Ana, Dave", "Ana, Dave, Jose", "Dave, Jose, Ana"]),
    ]
    lines = [["Job"] + [mdyyyy(x) for x in WEEK]] + [[j] + c for j, c in rows]
    exp = expected("C", "input.csv", layout="job_by_day", arrangement="single",
                   assignments=job_grid(rows, WEEK),
                   flags=[UF(S, "10-15"), NB("Luis", "10-15"), DB("Ana", "10-16", S, O)], info=[],
                   must_not_ask=NO_LAYOUT_QS)
    emit("C03", "job-by-day-multi-names",
         "Job x day grid with several first names per cell, split on ',', '&', '/', '+' and ' and '; an empty cell is an unstaffed job-day.",
         {"input.csv": csv_bytes(lines)}, exp, {"window": WEEK, "jobs": [J, S, O]})


def c04():
    cols = [
        (J, ["Mike Baker, Dave Kim"] * 3 + ["Mike Baker", "Mike Baker, Dave Kim"]),
        (S, ["Jose Garcia"] * 3 + ["Jose Garcia, Dave Kim", "Jose Garcia"]),
        (O, ["Luis Ortiz"] * 3 + ["", "Luis Ortiz"]),
        ("Off", ["", "", "", "Luis Ortiz", ""]),
    ]
    lines = [["Date"] + [c for c, _ in cols]]
    for i, x in enumerate(WEEK):
        lines.append([mdyyyy(x)] + [cells[i] for _, cells in cols])
    exp = expected("C", "input.csv", layout="day_by_job", arrangement="single",
                   assignments=job_grid(cols, WEEK),
                   flags=[UF(O, "10-15")], info=[], must_not_ask=NO_LAYOUT_QS)
    emit("C04", "day-by-job",
         "Day x job grid: dates down the side, jobs across, names in cells; an 'Off' column lists who is off and is never a job.",
         {"input.csv": csv_bytes(lines)}, exp, {"window": WEEK, "jobs": [J, S, O]})


def c05():
    people = [
        ("Mike Baker", [J] * 5),
        ("Jose Garcia", [S] * 5),
        ("Luis Ortiz", [O, O, O, O, "vac"]),
        ("Dave Kim", [J, J, f"{J} / {O}", O, O]),
    ]
    lines = [["Date"] + [p for p, _ in people]]
    for i, x in enumerate(WEEK):
        lines.append([mdyyyy(x)] + [cells[i] for _, cells in people])
    exp = expected("C", "input.csv", layout="day_by_person", arrangement="single",
                   assignments=person_grid(people, WEEK),
                   flags=[DB("Dave Kim", "10-14", J, O)], info=[], must_not_ask=NO_LAYOUT_QS)
    emit("C05", "day-by-person", "Day x person grid (the transpose of person x day): dates down the side, people across.",
         {"input.csv": csv_bytes(lines)}, exp, {"window": WEEK})


def c06():
    crews = {
        "Crew 1": ["Mike Baker", "Jose Garcia"],
        "Crew 2": ["Luis Ortiz", "Dave Kim", "Tom Reed"],
        "Crew 3": ["Ana Ruiz", "Carlos Mendez"],
    }
    plan = [("Crew 1", [J] * 5), ("Crew 2", [S, S, S, O, O]), ("Crew 3", [O, O, O, O, ""])]
    blank = [""] * 6
    lines = [["Crew"] + [mdyyyy(x) for x in WEEK]] + [[c] + cells for c, cells in plan]
    lines += [blank, ["Crew roster"] + [""] * 5]
    lines += [[c] + m + [""] * (5 - len(m)) for c, m in crews.items()]
    assigns = []
    for crew, cells in plan:
        for x, site in zip(WEEK, cells):
            if site:
                assigns += [W(m, x, site) for m in crews[crew]]
    exp = expected("C", "input.csv", layout="crew_based", arrangement="single",
                   assignments=assigns,
                   flags=[NB("Ana Ruiz", "10-16"), NB("Carlos Mendez", "10-16")], info=[],
                   must_not_ask=NO_LAYOUT_QS)
    roster = [m for ms in crews.values() for m in ms]
    emit("C06", "crew-based-roster-block",
         "Crew-based: crew names x days in the schedule, members in a separate crew roster block below; crews expand to members and never appear as people.",
         {"input.csv": csv_bytes(lines)}, exp, {"window": WEEK, "roster": roster})


def c07():
    w1 = [("Mike Baker", [J] * 5), ("Jose Garcia", [S] * 5), ("Luis Ortiz", [O] * 5), ("Dave Kim", [J] * 5)]
    w2 = [("Mike Baker", [J, J, S, S, S]), ("Jose Garcia", [S] * 5),
          ("Luis Ortiz", [O, O, O, f"{O} / {S}", O]), ("Dave Kim", [J] * 5)]
    blank = [""] * 6
    lines = [["Week of 10/12/2026"] + [""] * 5, ["Name"] + [mdyyyy(x) for x in WEEK]] + [[p] + c for p, c in w1]
    lines += [blank, ["Week of 10/19/2026"] + [""] * 5, ["Name"] + [mdyyyy(x) for x in WEEK2]] + [[p] + c for p, c in w2]
    exp = expected("C", "input.csv", layout="person_by_day", arrangement="stacked",
                   assignments=person_grid(w1, WEEK) + person_grid(w2, WEEK2),
                   flags=[DB("Luis Ortiz", "10-22", O, S)], info=[], must_not_ask=NO_LAYOUT_QS)
    emit("C07", "two-weeks-stacked",
         "Two person x day weeks stacked one under the other, each with its own 'Week of' title and header row.",
         {"input.csv": csv_bytes(lines)}, exp, {"window": WEEK + WEEK2})


def c08():
    w1 = [("Ana Ruiz", [S] * 5), ("Tom Reed", [S] * 5), ("Carlos Mendez", [P] * 5), ("Ray Jones", [J] * 5)]
    w2 = [("Ana Ruiz", [S] * 5), ("Tom Reed", [P] * 5), ("Carlos Mendez", [P, P, P, "", P]), ("Ray Jones", [J] * 5)]
    lines = [["Name"] + [mdyyyy(x) for x in WEEK] + ["", "Name"] + [mdyyyy(x) for x in WEEK2]]
    for (p1, c1), (p2, c2) in zip(w1, w2):
        lines.append([p1] + c1 + ["", p2] + c2)
    exp = expected("C", "input.csv", layout="person_by_day", arrangement="side_by_side",
                   assignments=person_grid(w1, WEEK) + person_grid(w2, WEEK2),
                   flags=[NB("Carlos Mendez", "10-22")], info=[], must_not_ask=NO_LAYOUT_QS)
    emit("C08", "two-weeks-side-by-side",
         "Two person x day weeks side by side, split by a blank column, with the Name column repeated for week two.",
         {"input.csv": csv_bytes(lines)}, exp, {"window": WEEK + WEEK2})


def c09():
    from openpyxl.styles import Alignment

    entries = {
        "10-12": [(J, ["Mike", "Jose"]), (S, ["Luis"]), (O, ["Dave"])],
        "10-13": [(J, ["Mike", "Jose"]), (S, ["Luis"]), (O, ["Dave"])],
        "10-14": [(J, ["Mike", "Jose"]), (S, ["Luis", "Mike"]), (O, ["Dave"])],
        "10-15": [(J, ["Jose"]), (S, ["Luis", "Mike"]), (O, ["Dave"])],
        "10-16": [(J, ["Jose"]), (S, ["Luis", "Mike"]), (O, ["Dave"])],
    }
    assigns = [W(n, day, site) for day, jobs in entries.items() for site, names in jobs for n in names]
    text = {d(day): "\n".join(f"{site}: {', '.join(names)}" for site, names in jobs) for day, jobs in entries.items()}

    def make(path):
        wb = new_workbook()
        ws = wb.active
        ws.title = "October"
        ws["A1"] = "October 2026"
        ws.merge_cells("A1:G1")
        for i, name in enumerate(["Sun", "Mon", "Tue", "Wed", "Thu", "Fri", "Sat"]):
            ws.cell(row=2, column=1 + i, value=name)
        cal = calendar.Calendar(firstweekday=6)  # Sunday first
        row = 3
        for week in cal.monthdatescalendar(2026, 10):
            for i, day in enumerate(week):
                if day.month == 10:
                    ws.cell(row=row, column=1 + i, value=day.day)
                    t = text.get(day.isoformat())
                    if t:
                        c = ws.cell(row=row + 1, column=1 + i, value=t)
                        c.alignment = Alignment(wrap_text=True, vertical="top")
            ws.row_dimensions[row + 1].height = 60
            row += 2
        for col in "ABCDEFG":
            ws.column_dimensions[col].width = 24
        wb.save(path)

    # weekday check for the calendar placement
    assert dt.date(2026, 10, 1).strftime("%a") == "Thu" and dt.date(2026, 10, 12).strftime("%a") == "Mon"
    exp = expected("C", "input.xlsx", layout="calendar_month", arrangement="single",
                   assignments=assigns, flags=[DB("Mike", "10-14", J, S)], info=[],
                   must_not_ask=[Q("pick_layout"), Q("pick_sheet")])
    emit("C09", "calendar-month-grid",
         "Calendar-month grid in .xlsx: 'October 2026', Sun-Sat columns, a day-number row over an entries row per week; entries are 'Site: names' lines.",
         {"input.xlsx": make}, exp)


def c10():
    rows = [("Mike Baker", [J] * 5), ("Jose Garcia", [S] * 5), ("Luis Ortiz", [O] * 5), ("Dave Kim", [J, J, J, S, S])]
    roster = [("Mike Baker", "555-0101"), ("Jose Garcia", "555-0102"), ("Luis Ortiz", "555-0103"),
              ("Dave Kim", "555-0104"), ("Kevin Walsh", "555-0105")]
    lines = [["Name"] + [mdyyyy(x) for x in WEEK] + ["", "Roster", "Phone"]]
    for i, (name, phone) in enumerate(roster):
        left = [rows[i][0]] + rows[i][1] if i < len(rows) else [""] * 6
        lines.append(left + ["", name, phone])
    exp = expected("C", "input.csv", layout="person_by_day", arrangement="single",
                   assignments=person_grid(rows, WEEK),
                   flags=[NB("Kevin Walsh", x) for x in WEEK], info=[], must_not_ask=NO_LAYOUT_QS)
    emit("C10", "schedule-and-roster-one-sheet",
         "Schedule and roster on one sheet, side by side across a blank column; the roster adds Kevin Walsh, who is on no job all week.",
         {"input.csv": csv_bytes(lines)}, exp, {"window": WEEK, "roster": [n for n, _ in roster]})


def c11():
    rows = [("Mike Baker", [J] * 5), ("Jose Garcia", [S] * 5), ("Luis Ortiz", [O, O, f"{O} / {S}", O, O]),
            ("Dave Kim", [J] * 5), ("Ana Ruiz", [O] * 5)]

    def make(path):
        wb = new_workbook()
        ws = wb.active
        ws.title = "Contacts"
        ws.append(["Name", "Phone", "Role"])
        for name, phone, role in [("Mike Baker", "555-0101", "Lead carpenter"), ("Jose Garcia", "555-0102", "Carpenter"),
                                  ("Luis Ortiz", "555-0103", "Tile setter"), ("Dave Kim", "555-0104", "Framer"),
                                  ("Ana Ruiz", "555-0106", "Carpenter")]:
            ws.append([name, phone, role])
        write_person_sheet(wb.create_sheet("Schedule"), rows, WEEK)
        m = wb.create_sheet("Materials")
        m.append(["Date ordered", "Supplier", "Item", "Qty"])
        for day, sup, item, qty in [("10-05", "Home Depot", "2x4x8 studs", 120), ("10-06", "ABC Supply", "Tyvek wrap", 3),
                                    ("10-07", "84 Lumber", "Deck screws 3in", 10), ("10-08", "Home Depot", "Joint compound", 6)]:
            m.append([pydate(d(day)), sup, item, qty])
            m.cell(row=m.max_row, column=1).number_format = "m/d/yyyy"
        n = wb.create_sheet("Notes")
        for line in ["Order more studs before the framing inspection.", "Trailer needs new tires.",
                     "Ask the city about the permit extension."]:
            n.append([line])
        wb.save(path)

    exp = expected("C", "input.xlsx", layout="person_by_day", arrangement="single",
                   assignments=person_grid(rows, WEEK),
                   flags=[DB("Luis Ortiz", "10-14", O, S)], info=[],
                   must_not_ask=NO_LAYOUT_QS + [Q("pick_sheet")])
    emit("C11", "workbook-one-schedule-sheet",
         "Workbook with four sheets (Contacts first, then Schedule, Materials, Notes); only 'Schedule' is a schedule and it must be picked without asking.",
         {"input.xlsx": make}, exp, {"window": WEEK})


def c12():
    w1 = [("Mike Baker", [J] * 5), ("Jose Garcia", [S] * 5), ("Luis Ortiz", [O] * 5), ("Dave Kim", [J, J, f"{J} / {S}", S, S])]
    w2 = [("Mike Baker", [S] * 5), ("Jose Garcia", [S] * 5), ("Luis Ortiz", [O] * 5), ("Dave Kim", [J] * 5)]

    def make(path):
        wb = new_workbook()
        ws = wb.active
        ws.title = "Week 10-12"
        write_person_sheet(ws, w1, WEEK)
        write_person_sheet(wb.create_sheet("Week 10-19"), w2, WEEK2)
        pl = wb.create_sheet("Phone list")
        pl.append(["Name", "Phone"])
        for name, phone in [("Mike Baker", "555-0101"), ("Jose Garcia", "555-0102"), ("Luis Ortiz", "555-0103"), ("Dave Kim", "555-0104")]:
            pl.append([name, phone])
        wb.save(path)

    q = Q("pick_sheet", options=["Week 10-12", "Week 10-19"])
    exp = expected("C", "input.xlsx", layout="person_by_day", arrangement="single",
                   ask=[q], must_not_ask=[Q("pick_layout"), Q("column_role")],
                   answers=[{"question": Q("pick_sheet"), "answer": "Week 10-12"}],
                   after_answers={"assignments": person_grid(w1, WEEK),
                                  "flags": [DB("Dave Kim", "10-14", J, S)], "info": []})
    emit("C12", "workbook-two-candidate-sheets",
         "Workbook with two schedule-looking sheets ('Week 10-12', 'Week 10-19') plus a phone list: must ask pick_sheet; answering 'Week 10-12' reads that sheet.",
         {"input.xlsx": make}, exp, {"window": WEEK})


def c13():
    L, R, E = "Lincoln School", "Riverside duplex", "88 Elm Ave"
    rows = [
        (L, ["Juan Lopez, Miguel Torres"] * 3 + ["Juan Lopez", "Juan Lopez, Miguel Torres"]),
        (R, ["Chris Hall", "Chris Hall", "Chris Hall, Brian Young", "Chris Hall, Miguel Torres", ""]),
        (E, ["Brian Young", "Brian Young", "Eddie Flores", "Brian Young", "Brian Young, Chris Hall"]),
        ("Shop", ["Eddie Flores", "Eddie Flores", "", "Eddie Flores", "Eddie Flores"]),
    ]

    def make(path):
        from odf.opendocument import OpenDocumentSpreadsheet
        from odf.table import Table, TableCell, TableRow
        from odf.text import P as OP

        doc = OpenDocumentSpreadsheet()
        t = Table(name="Schedule")

        def text_cell(s):
            c = TableCell(valuetype="string")
            if s:
                c.addElement(OP(text=s))
            else:
                c = TableCell()
            return c

        hdr = TableRow()
        hdr.addElement(text_cell("Job"))
        for x in WEEK:
            c = TableCell(valuetype="date", datevalue=x)
            c.addElement(OP(text=mdyyyy(x)))
            hdr.addElement(c)
        t.addElement(hdr)
        for job, cells in rows:
            tr = TableRow()
            tr.addElement(text_cell(job))
            for cell in cells:
                tr.addElement(text_cell(cell))
            t.addElement(tr)
        doc.spreadsheet.addElement(t)
        doc.save(str(path))

    exp = expected("C", "input.ods", layout="job_by_day", arrangement="single",
                   assignments=job_grid(rows, WEEK), flags=[UF(R, "10-16")], info=[],
                   must_not_ask=NO_LAYOUT_QS + [Q("pick_sheet")])
    emit("C13", "ods-job-by-day",
         "OpenDocument spreadsheet (.ods, LibreOffice) holding a job x day grid with date-typed header cells; the empty Shop cell is never unstaffed.",
         {"input.ods": make}, exp, {"window": WEEK, "jobs": [L, R, E, "Shop"]})


def c14():
    rows = [("Mike Baker", [J] * 5), ("Jose Garcia", [S] * 5), ("Luis Ortiz", [O, O, O, O, "PTO"]), ("Ana Ruiz", [S] * 5)]

    def make(path):
        from docx import Document

        doc = Document()
        cp = doc.core_properties
        cp.author = cp.last_modified_by = "Office"
        cp.created = cp.modified = FIXED
        cp.revision = 1
        doc.add_paragraph("Crew Schedule — Week of 10/12/2026")
        table = doc.add_table(rows=1 + len(rows), cols=6)
        table.style = "Table Grid"
        for i, h in enumerate(["Name"] + [wd_md(x) for x in WEEK]):
            table.cell(0, i).text = h
        for r, (person, cells) in enumerate(rows, start=1):
            table.cell(r, 0).text = person
            for i, cell in enumerate(cells):
                table.cell(r, 1 + i).text = cell
        doc.add_paragraph("")
        doc.add_paragraph("Call Dave with any changes.")
        doc.save(str(path))

    exp = expected("C", "input.docx", layout="person_by_day", arrangement="single",
                   assignments=person_grid(rows, WEEK), flags=[], info=[],
                   must_not_ask=NO_LAYOUT_QS + [Q("weekday_mismatch")])
    emit("C14", "docx-word-table",
         "A person x day table typed in Word (.docx), with a title paragraph above and a 'Call Dave' line below that are not data.",
         {"input.docx": make}, exp, {"window": WEEK})


if __name__ == "__main__":
    for f in (c01, c02, c03, c04, c05, c06, c07, c08, c09, c10, c11, c12, c13, c14):
        f()
    print("C done")
