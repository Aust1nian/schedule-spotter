"""Category I: junk and noise. Run: uv run --with openpyxl python3 build_I.py"""
import datetime as dt
import re
import zipfile
from xml.sax.saxutils import escape

from fxlib import *

WEEK = days("10-12", 5)
J, S, O = "1420 Johnson St", "Smith remodel", "Oak Ridge Apts"
P, E = "Parker kitchen", "88 Elm Ave"
NO_LAYOUT_QS = [Q("pick_layout"), Q("column_role")]
FIXED = dt.datetime(2026, 10, 9, 12, 0)


def i01():
    rows = [("Mike Baker", [J] * 5), ("Jose Garcia", [S] * 5), ("Luis Ortiz", [O, O, O, O, "off"]),
            ("Dave Kim", [J, J, f"{J} / {S}", J, J])]
    worked = [sum(1 for c in cells if c and c != "off") for _, cells in rows]
    per_day = [sum(1 for _, cells in rows if cells[i] and cells[i] != "off") for i in range(5)]
    assert worked == [5, 5, 4, 5] and per_day == [4, 4, 4, 4, 3]
    lines = [["Name"] + [mdyyyy(x) for x in WEEK] + ["Days"]]
    lines += [[p] + c + [str(n)] for (p, c), n in zip(rows, worked)]
    lines += [["Total"] + [str(n) for n in per_day] + [str(sum(per_day))]]
    exp = expected("I", "input.csv", layout="person_by_day", arrangement="single",
                   assignments=person_grid(rows, WEEK),
                   flags=[DB("Dave Kim", "10-14", J, S)], info=[], must_not_ask=NO_LAYOUT_QS)
    emit("I01", "totals-rows",
         "A 'Total' headcount row under the grid and a 'Days' count column at its right; numbers are not sites and 'Total' is not a person.",
         {"input.csv": csv_bytes(lines)}, exp, {"window": WEEK})


def i02():
    rows = [("Ana Ruiz", [P] * 5), ("Tom Reed", [P] * 5), ("Carlos Mendez", [J] * 5), ("Ray Jones", [J, J, J, J, "PTO"])]
    lines = [["Date", "Name", "Job"]] + grid_to_list_rows(rows, WEEK, mdyyyy)
    lines += [
        ["", "", ""],
        ["Notes:", "", ""],
        ["Johnson inspection Wed 9am - call city", "", ""],
        ["Bring the trailer back to the yard Friday", "", ""],
        ["", "", ""],
        ["Approved by: Rick Hansen (owner)", "", ""],
        ["Signature: ________________   Date: __________", "", ""],
    ]
    exp = expected("I", "input.csv", layout="long_list", arrangement="single",
                   assignments=person_grid(rows, WEEK), flags=[], info=[], must_not_ask=NO_LAYOUT_QS)
    emit("I02", "footer-notes-signature",
         "Footer notes (with a weekday, a time and a site name) and approval/signature lines under the schedule; none are data and Rick Hansen is not on the roster.",
         {"input.csv": csv_bytes(lines, eol="\r\n")}, exp)


def i03():
    rows = [("Mike Baker", [J] * 5), ("Jose Garcia", [S] * 5), ("Luis Ortiz", [O, O, O, O, f"{O} / {S}"]),
            ("Dave Kim", [J, J, "off", J, J])]

    def spread(first, cells):  # leading blank column, a blank column between Wed and Thu, a trailing blank
        return ["", first] + cells[:3] + [""] + cells[3:] + [""]

    blank = [""] * 9
    lines = [blank, spread("Name", [mdyyyy(x) for x in WEEK]), blank]
    for i, (p, c) in enumerate(rows):
        lines.append(spread(p, c))
        if i == 1:
            lines += [blank, blank]
    lines.append(blank)
    exp = expected("I", "input.csv", layout="person_by_day", arrangement="single",
                   assignments=person_grid(rows, WEEK),
                   flags=[DB("Luis Ortiz", "10-16", O, S)], info=[],
                   must_not_ask=NO_LAYOUT_QS + [Q("fill_down")])
    emit("I03", "empty-rows-and-columns",
         "Empty rows and columns sprinkled through one grid (leading column, a column between Wed and Thu, rows between people); still one single grid.",
         {"input.csv": csv_bytes(lines)}, exp, {"window": WEEK})


def i04():
    groups = [
        (J, [("10-12", "Mike Baker"), ("10-12", "Dave Kim"), ("10-13", "Mike Baker"), ("10-13", "Dave Kim"),
             ("10-14", "Mike Baker"), ("10-14", "Dave Kim"), ("10-15", "Mike Baker"), ("10-15", "Dave Kim"),
             ("10-16", "Mike Baker"), ("10-16", "Dave Kim")]),
        (S, [("10-12", "Jose Garcia"), ("10-13", "Jose Garcia"), ("10-14", "Jose Garcia"), ("10-14", "Luis Ortiz"),
             ("10-15", "Jose Garcia"), ("10-16", "Jose Garcia")]),
        (O, [(x, "Luis Ortiz") for x in ("10-12", "10-13", "10-14", "10-15", "10-16")]),
    ]
    assigns = [W(name, day, job) for job, items in groups for day, name in items]

    def make(path):
        import openpyxl
        from openpyxl.styles import Alignment
        wb = openpyxl.Workbook()
        wb.properties.creator = "Office"
        wb.properties.created = wb.properties.modified = FIXED
        ws = wb.active
        ws.title = "Schedule"
        ws.append(["Job", "Date", "Name"])
        row = 2
        for job, items in groups:
            first = row
            for day, name in items:
                c = ws.cell(row=row, column=2, value=pydate(d(day)))
                c.number_format = "m/d/yyyy"
                ws.cell(row=row, column=3, value=name)
                row += 1
            ws.cell(row=first, column=1, value=job).alignment = Alignment(vertical="center")
            ws.merge_cells(start_row=first, start_column=1, end_row=row - 1, end_column=1)
        wb.save(path)

    exp = expected("I", "input.xlsx", layout="long_list", arrangement="single",
                   assignments=assigns, flags=[DB("Luis Ortiz", "10-14", S, O)], info=[],
                   must_not_ask=NO_LAYOUT_QS + [Q("fill_down"), Q("pick_sheet")])
    emit("I04", "merged-cell-blanks-xlsx",
         "Job written once in a vertically merged .xlsx cell over all of its rows; the merge must fill down without asking.",
         {"input.xlsx": make}, exp)


def i05():
    groups = [
        (P, [(x, n) for x in WEEK for n in ("Ana Ruiz", "Tom Reed")]),
        (E, [(x, n) for x in WEEK for n in ("Carlos Mendez", "Ray Jones") if not (n == "Ray Jones" and x == "2026-10-16")]),
    ]
    lines = [["Job", "Date", "Name"]]
    assigns = []
    for job, items in groups:
        for i, (x, name) in enumerate(items):
            lines.append([job if i == 0 else "", mdyyyy(x), name])
            assigns.append(W(name, x, job))
    exp = expected("I", "input.csv", layout="long_list", arrangement="single",
                   assignments=assigns, flags=[NB("Ray Jones", "10-16")], info=[],
                   must_not_ask=NO_LAYOUT_QS + [Q("fill_down")])
    emit("I05", "merged-cell-pattern-csv",
         "The merged-cell pattern after CSV export: the job is written on the first row of each group with blanks below, every other cell filled; fill down without asking.",
         {"input.csv": csv_bytes(lines)}, exp)


# ------------------------------------------------------------------ I06: hand-built xlsx with cached formula results

def col(n):
    s = ""
    while n:
        n, r = divmod(n - 1, 26)
        s = chr(65 + r) + s
    return s


def i06():
    LOOKUP = [("J-1042", J), ("J-1043", S), ("J-1044", O)]
    vl = 'VLOOKUP("{}",Lookup!$A$2:$B$4,2,FALSE)'
    # each cell: ("s", text) plain string | ("f", formula, cached text) | ("e", formula, error)
    grid = [
        ("Mike Baker", [("s", J), ("e", "#REF!", "#REF!"), ("s", J), ("s", J), ("s", J)]),
        ("Jose Garcia", [("s", S), ("s", S), ("s", S), ("e", vl.format("J-1050"), "#N/A"), ("s", S)]),
        ("Luis Ortiz", [("s", O), ("f", "B4", O), ("f", "C4", O), ("f", "D4", O), ("f", "E4", O)]),
        ("Dave Kim", [("s", J)] + [("f", vl.format("J-1042"), J)] * 4),
        ("Ana Ruiz", [("s", S)] * 5),
    ]
    rows = [(p, ["" if c[0] == "e" else c[-1] for c in cells]) for p, cells in grid]
    assigns = person_grid(rows, WEEK)

    def make(path):
        strings = []

        def sst(t):
            if t not in strings:
                strings.append(t)
            return strings.index(t)

        def sheet_xml(rows_xml):
            return ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
                    '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
                    f'<sheetData>{"".join(rows_xml)}</sheetData></worksheet>')

        r1 = [f'<c r="A1" t="s"><v>{sst("Name")}</v></c>']
        for i, x in enumerate(WEEK):
            r1.append(f'<c r="{col(2 + i)}1" s="1"><v>{serial(x)}</v></c>')
        rows_xml = [f'<row r="1">{"".join(r1)}</row>']
        for r, (person, cells) in enumerate(grid, start=2):
            cs = [f'<c r="A{r}" t="s"><v>{sst(person)}</v></c>']
            for i, c in enumerate(cells):
                ref = f"{col(2 + i)}{r}"
                if c[0] == "s":
                    cs.append(f'<c r="{ref}" t="s"><v>{sst(c[1])}</v></c>')
                elif c[0] == "f":
                    cs.append(f'<c r="{ref}" t="str"><f>{escape(c[1])}</f><v>{escape(c[2])}</v></c>')
                else:
                    cs.append(f'<c r="{ref}" t="e"><f>{escape(c[1])}</f><v>{escape(c[2])}</v></c>')
            rows_xml.append(f'<row r="{r}">{"".join(cs)}</row>')
        sheet1 = sheet_xml(rows_xml)

        lk = [f'<row r="1"><c r="A1" t="s"><v>{sst("Job #")}</v></c><c r="B1" t="s"><v>{sst("Job name")}</v></c></row>']
        for r, (code, name) in enumerate(LOOKUP, start=2):
            lk.append(f'<row r="{r}"><c r="A{r}" t="s"><v>{sst(code)}</v></c><c r="B{r}" t="s"><v>{sst(name)}</v></c></row>')
        sheet2 = sheet_xml(lk)

        ss = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
              f'<sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" count="{len(strings)}" uniqueCount="{len(strings)}">'
              + "".join(f"<si><t>{escape(t)}</t></si>" for t in strings) + "</sst>")
        parts = {
            "[Content_Types].xml":
                '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
                '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
                '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
                '<Default Extension="xml" ContentType="application/xml"/>'
                '<Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>'
                '<Override PartName="/xl/worksheets/sheet1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>'
                '<Override PartName="/xl/worksheets/sheet2.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>'
                '<Override PartName="/xl/styles.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.styles+xml"/>'
                '<Override PartName="/xl/sharedStrings.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml"/>'
                '</Types>',
            "_rels/.rels":
                '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
                '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="xl/workbook.xml"/>'
                '</Relationships>',
            "xl/workbook.xml":
                '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
                '<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
                'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
                '<sheets><sheet name="Schedule" sheetId="1" r:id="rId1"/><sheet name="Lookup" sheetId="2" r:id="rId2"/></sheets>'
                '<calcPr calcId="191029"/></workbook>',
            "xl/_rels/workbook.xml.rels":
                '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
                '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet1.xml"/>'
                '<Relationship Id="rId2" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet2.xml"/>'
                '<Relationship Id="rId3" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/styles" Target="styles.xml"/>'
                '<Relationship Id="rId4" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/sharedStrings" Target="sharedStrings.xml"/>'
                '</Relationships>',
            "xl/styles.xml":
                '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
                '<styleSheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
                '<fonts count="1"><font><sz val="11"/><name val="Calibri"/></font></fonts>'
                '<fills count="2"><fill><patternFill patternType="none"/></fill><fill><patternFill patternType="gray125"/></fill></fills>'
                '<borders count="1"><border><left/><right/><top/><bottom/><diagonal/></border></borders>'
                '<cellStyleXfs count="1"><xf numFmtId="0" fontId="0" fillId="0" borderId="0"/></cellStyleXfs>'
                '<cellXfs count="2"><xf numFmtId="0" fontId="0" fillId="0" borderId="0" xfId="0"/>'
                '<xf numFmtId="14" fontId="0" fillId="0" borderId="0" xfId="0" applyNumberFormat="1"/></cellXfs>'
                '<cellStyles count="1"><cellStyle name="Normal" xfId="0" builtinId="0"/></cellStyles>'
                '</styleSheet>',
            "xl/worksheets/sheet1.xml": sheet1,
            "xl/worksheets/sheet2.xml": sheet2,
            "xl/sharedStrings.xml": ss,
        }
        with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
            for name, text in parts.items():
                z.writestr(zipfile.ZipInfo(name, date_time=FIXED_TIME), text.encode("utf-8"), compress_type=zipfile.ZIP_DEFLATED)

    exp = expected("I", "input.xlsx", layout="person_by_day", arrangement="single",
                   assignments=assigns,
                   flags=[NB("Mike Baker", "10-13"), NB("Jose Garcia", "10-15")], info=[],
                   must_not_ask=NO_LAYOUT_QS + [Q("pick_sheet")],
                   notes=["#REF!", "#N/A"])
    emit("I06", "error-values-and-formulas",
         "Real formulas in .xlsx: a cached #REF! and a cached #N/A (both read as empty, with a note) next to formulas whose cached text values must be used, never the formula text.",
         {"input.xlsx": make}, exp, {"window": WEEK})


def i07():
    junk_headers = ["Emp ID", "Truck", "Cost code", "Phase", "Hrs", "OT hrs", "Rate", "Per diem", "Miles", "Union",
                    "Class", "Dept", "Region", "Badge", "Radio", "Fuel card", "PO", "Budget code", "Tax code",
                    "WC code", "Crew size", "Equip 1", "Equip 2", "Equip 3", "Permit no", "Lot", "Block", "Zone",
                    "Bid item", "Change order", "Retainage", "Billing", "Approved"] + [f"Ref {i}" for i in range(1, 15)]
    header = ["Emp ID", "Date", "Employee", "Truck", "Job"] + junk_headers[2:]
    assert len(header) == 50 and len(set(header)) == 50
    who = {"Mike Baker": ("1042", "12"), "Jose Garcia": ("1057", "7"), "Luis Ortiz": ("1063", "7"), "Dave Kim": ("1071", "3")}
    rows = [("Mike Baker", [J] * 5), ("Jose Garcia", [S] * 5), ("Luis Ortiz", [S] * 5), ("Dave Kim", [J, J, J, J, "off"])]
    lines = [header]
    for date_text, person, site in grid_to_list_rows(rows, WEEK, mdyyyy):
        emp, truck = who[person]
        working = site != "off"
        extra = {
            "Cost code": "06100" if working else "", "Phase": ("Rough" if site == J else "Finish") if working else "",
            "Hrs": "8" if working else "0", "OT hrs": "0", "Rate": "32.50", "Per diem": "0", "Miles": "14" if working else "0",
            "Union": "Y", "Class": "C2", "Dept": "Field", "Region": "North", "Badge": "B" + emp, "Radio": "Ch 3",
            "Fuel card": "4471", "PO": "PO 5521" if working else "", "Budget code": "B200", "Tax code": "TX1",
            "WC code": "5645", "Crew size": "2" if working else "", "Equip 1": "Compressor" if site == J else "",
            "Lot": "7" if site == S else "", "Block": "C" if site == S else "", "Retainage": "10%", "Billing": "TM",
            "Approved": "Y",
        }
        lines.append([emp, date_text, person, truck, site] + [extra.get(h, "") for h in header[5:]])
    exp = expected("I", "input.csv", layout="long_list", arrangement="single",
                   assignments=person_grid(rows, WEEK), flags=[], info=[], must_not_ask=NO_LAYOUT_QS)
    emit("I07", "fifty-columns-wide",
         "A 50-column payroll-style export where only Date, Employee and Job matter; the other 47 columns (codes, numbers, blanks) are ignored.",
         {"input.csv": csv_bytes(lines)}, exp)


LONG_NOTE = " ".join([
    "Gate code is 4471.",
    "Park along the curb past the second hydrant so the pump truck can reach the slab.",
    "Sign in at the trailer every morning before work begins.",
    "Hard hats vests glasses gloves are required everywhere on site.",
    "The east stair has no handrail yet so use the west stair until the inspector signs it.",
    "Keep the porta john door facing the street because the neighbor complained last week.",
    "Do not block the driveway of the blue house at the corner.",
    "The homeowner works nights so keep the radio low near the bedroom wall.",
    "Bring the long ladder from the trailer for the soffit work on the north side.",
    "The window order is short two units so frame the openings but leave them covered with plastic.",
    "Check every joist hanger on the second floor deck before the inspector comes back.",
    "Nails for the hangers are in the gray bucket under the stairs.",
    "If it rains cover the subfloor with the tarps stacked by the dumpster.",
    "The dumpster swap is booked through the rental company so leave the gate unlocked for the driver.",
    "Sweep the deck at the end of each day because the framing inspector failed us last time for debris.",
    "The electrician will be roughing in the kitchen so stay clear of that room while he pulls wire.",
    "Measure twice for the stair stringers because the landing height changed on the revised plans.",
    "The revised plans are in the red folder in the trailer.",
    "Do not cut any engineered beam without calling Rob first.",
    "Rob is the super for this job. His cell is 555 0199.",
    "Lunch trucks come by around noon on the main road.",
    "Water is in the cooler by the trailer steps. Refill it before you leave.",
    "Lock the trailer every night. The lockbox code is 2218.",
    "Return the generator to the truck when the power company finishes the meter.",
    "Take photos of every wall before insulation goes in.",
    "Upload the photos to the job folder the same day.",
    "The garage slab gets poured next week so keep the forms clear of tools.",
    "Ask Rob before moving any of the shoring posts in the basement.",
    "Shut every window at night to keep rain from the new subfloor.",
    "Questions go to Rob first. Then call the main number.",
])
BANNED = re.compile(
    r"[,/&+()\-?:;\"\n]|\b(and|y|off|out|sick|x|holiday|vacation|vac|pto|open|tbd|need|am|pm|half|shop|yard|office|"
    r"warehouse|closed|done|monday|tuesday|wednesday|thursday|friday|saturday|sunday|mon|tue|wed|thu|fri|sat|sun)\b", re.I)


def i08():
    assert len(LONG_NOTE) >= 2000, len(LONG_NOTE)
    assert not BANNED.search(LONG_NOTE), BANNED.search(LONG_NOTE)
    rows = [("Mike Baker", [J] * 5), ("Jose Garcia", [S] * 5),
            ("Luis Ortiz", [f"{O} ({LONG_NOTE})", O, O, O, O]), ("Dave Kim", [S] * 5)]
    lines = [["Name"] + [mdyyyy(x) for x in WEEK]] + [[p] + c for p, c in rows]
    assigns = person_grid(rows, WEEK)
    assert assigns[10]["note"] == LONG_NOTE
    exp = expected("I", "input.csv", layout="person_by_day", arrangement="single",
                   assignments=assigns, flags=[], info=[], must_not_ask=NO_LAYOUT_QS)
    emit("I08", "very-long-cell",
         f"One grid cell holds a site plus a {len(LONG_NOTE)}-character parenthetical note (site instructions); the note goes to `note` whole, nothing else changes.",
         {"input.csv": csv_bytes(lines)}, exp, {"window": WEEK})


def i09():
    S2 = "Smith remodel \U0001F3E0"
    worker = "\U0001F477‍♂️"  # man construction worker, a ZWJ sequence
    rows = [
        ("Mike Baker", [f"{J} (\U0001F6A7 park on 15th)", J, J, J, J]),
        ("Jose Garcia", [S2] * 5),
        ("Dave Kim \U0001F527", [J, J, f"{J} / {S2}", J, J]),
        ("Luis Ortiz", [f"{O} ({worker} first day)", O, O, O, O]),
    ]
    lines = [["Name \U0001F477"] + [mdyyyy(x) for x in WEEK]] + [[p] + c for p, c in rows]
    exp = expected("I", "input.csv", layout="person_by_day", arrangement="single",
                   assignments=person_grid(rows, WEEK),
                   flags=[DB("Dave Kim \U0001F527", "10-14", J, S2)], info=[], must_not_ask=NO_LAYOUT_QS)
    emit("I09", "emoji-in-names-and-notes",
         "Emoji in a header, a person name ('Dave Kim 🔧'), a site name ('Smith remodel 🏠') and notes, including a ZWJ sequence; all kept as written.",
         {"input.csv": csv_bytes(lines)}, exp, {"window": WEEK})


if __name__ == "__main__":
    for f in (i01, i02, i03, i04, i05, i06, i07, i08, i09):
        f()
    print("I done")
