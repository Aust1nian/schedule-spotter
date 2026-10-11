"""Format coverage fixtures C15–C19: the same small schedules written as
Apple Numbers (.numbers), Excel 97-2003 (.xls), macro-enabled Excel (.xlsm)
and Excel binary (.xlsb), so every accepted extension has a fixture.

Run: uv run --with numbers-parser --with openpyxl --with xlwt \
       python3 fixtures/handmade/_make/build_formats.py

.xlsb has no common writer, so this file carries a minimal BIFF12 writer
(MS-XLSB records: BrtBundleSh, BrtWsDim, BrtRowHdr, BrtCellSt, BrtCellReal).
"""

import io
import json
import struct
import zipfile
from datetime import datetime
from pathlib import Path

HERE = Path(__file__).resolve().parent.parent
WORKDAYS = ["mon", "tue", "wed", "thu", "fri"]
DATES = ["2026-10-12", "2026-10-13", "2026-10-14", "2026-10-15", "2026-10-16"]

# person × day: one person's row per line.
PERSON_GRID = {
    "Mike Baker": ["1420 Johnson St", "1420 Johnson St", "1420 Johnson St / Smith remodel", "Smith remodel", "Smith remodel"],
    "José García": ["Smith remodel", "Smith remodel", "OFF", "Smith remodel", "Smith remodel"],
    "Luis Ortiz": ["Oak Ridge Apts", "Oak Ridge Apts", "Oak Ridge Apts", "PTO", "Oak Ridge Apts"],
    "Dave Kim": ["Shop", "1420 Johnson St", "1420 Johnson St", "1420 Johnson St", ""],
}
STATUS = {"OFF": "off", "PTO": "pto"}


def person_assignments(grid):
    out = []
    for person, cells in grid.items():
        for date, cell in zip(DATES, cells):
            if not cell:
                continue
            if cell in STATUS:
                out.append({"person": person, "date": date, "status": STATUS[cell]})
                continue
            for site in cell.split(" / "):
                out.append({"person": person, "date": date, "site": site})
    return out


INTERNAL = {"Shop"}


def person_flags(grid):
    flags = []
    staffed = {}  # job -> dates with someone on it
    for person, cells in grid.items():
        for date, cell in zip(DATES, cells):
            if " / " in cell:
                flags.append({"type": "double_booked", "person": person, "date": date, "sites": cell.split(" / ")})
            if not cell:
                flags.append({"type": "not_booked", "person": person, "date": date})
            if cell and cell not in STATUS:
                for site in cell.split(" / "):
                    if site not in INTERNAL:
                        staffed.setdefault(site, set()).add(date)
    # Jobs run from their first to last staffed working day (one week here).
    for job, days in staffed.items():
        first, last = min(days), max(days)
        for d in DATES:
            if first <= d <= last and d not in days:
                flags.append({"type": "unstaffed", "site": job, "date": d})
    return flags


def expected(category, input_name, layout, assignments, flags, extra=None):
    e = {
        "category": category,
        "tier": "script",
        "input": input_name,
        "paste": False,
        "settings": {"today": "2026-10-10", "workdays": WORKDAYS, "days_off": []},
        "layout": layout,
        "arrangement": "single",
        "assignments": assignments,
        "assignments_exact": True,
        "flags": flags,
        "flags_exact": True,
    }
    e.update(extra or {})
    return e


def write_fixture(slug, readme, input_name, data, exp):
    d = HERE / slug
    d.mkdir(exist_ok=True)
    (d / "README.md").write_text(readme + "\n")
    (d / input_name).write_bytes(data)
    (d / "expected.json").write_text(json.dumps(exp, indent=2, ensure_ascii=False) + "\n")


# ------------------------------------------------------------------ Numbers

def numbers_bytes(rows):
    from numbers_parser import Document

    doc = Document(num_header_rows=1, num_header_cols=1, num_rows=len(rows) + 2, num_cols=len(rows[0]) + 1)
    table = doc.sheets[0].tables[0]
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            if value not in ("", None):
                table.write(r, c, value)
    path = Path("/tmp/_spotter_fixture.numbers")
    doc.save(str(path))
    data = path.read_bytes()
    path.unlink()
    return data


def build_numbers():
    header = ["Name"] + [datetime.fromisoformat(d) for d in DATES]
    rows = [header] + [[p] + cells for p, cells in PERSON_GRID.items()]
    write_fixture(
        "C15-numbers-person-by-day",
        "Apple Numbers (.numbers) person × day grid with real date cells in the header, an accented name and statuses.",
        "input.numbers",
        numbers_bytes(rows),
        expected("C", "input.numbers", "person_by_day", person_assignments(PERSON_GRID), person_flags(PERSON_GRID)),
    )

    jobs = {
        "1420 Johnson St": ["Mike, Dave", "Mike, Dave", "Dave", "TBD", "Dave"],
        "Smith remodel": ["José", "José & Luis", "Mike, José", "José", ""],
    }
    rows = [["Job", "Mon 10/12", "Tue 10/13", "Wed 10/14", "Thu 10/15", "Fri 10/16"]]
    rows += [[job] + cells for job, cells in jobs.items()]
    assignments, flags = [], []
    for job, cells in jobs.items():
        for date, cell in zip(DATES, cells):
            if cell == "TBD":
                assignments.append({"date": date, "site": job, "unstaffed": True})
                flags.append({"type": "unstaffed", "site": job, "date": date})
            elif not cell:
                flags.append({"type": "unstaffed", "site": job, "date": date})
            else:
                for name in cell.replace(" & ", ", ").split(", "):
                    assignments.append({"person": name, "date": date, "site": job})
    # Everyone is on the roster; a weekday without them is "not booked".
    seen = {}
    for a in assignments:
        if "person" in a:
            seen.setdefault(a["person"], set()).add(a["date"])
    for person, days in seen.items():
        for d in DATES:
            if d not in days:
                flags.append({"type": "not_booked", "person": person, "date": d})
    write_fixture(
        "C16-numbers-job-by-day",
        "Apple Numbers (.numbers) job × day grid with several names per cell, a TBD and an empty job-day.",
        "input.numbers",
        numbers_bytes(rows),
        expected("C", "input.numbers", "job_by_day", assignments, flags),
    )


# ---------------------------------------------------------------------- xls

def build_xls():
    import xlwt

    rows = [
        ("Date", "Employee", "Job"),
        ("10/12/2026", "Mike Baker", "1420 Johnson St"),
        ("10/12/2026", "Jose Garcia", "Smith remodel"),
        ("10/13/2026", "Mike Baker", "1420 Johnson St"),
        ("10/13/2026", "Mike Baker", "Smith remodel"),
        ("10/13/2026", "Jose Garcia", "Smith remodel"),
        ("10/14/2026", "Jose Garcia", "PTO"),
        ("10/14/2026", "Mike Baker", "Smith remodel"),
    ]
    wb = xlwt.Workbook()
    ws = wb.add_sheet("Schedule")
    date_style = xlwt.easyxf(num_format_str="M/D/YYYY")
    for r, row in enumerate(rows):
        for c, v in enumerate(row):
            if r > 0 and c == 0:
                ws.write(r, c, datetime.strptime(v, "%m/%d/%Y"), date_style)
            else:
                ws.write(r, c, v)
    buf = io.BytesIO()
    wb.save(buf)
    assignments = []
    for d, p, j in rows[1:]:
        iso = datetime.strptime(d, "%m/%d/%Y").date().isoformat()
        if j == "PTO":
            assignments.append({"person": p, "date": iso, "status": "pto"})
        else:
            assignments.append({"person": p, "date": iso, "site": j})
    flags = [
        {"type": "double_booked", "person": "Mike Baker", "date": "2026-10-13", "sites": ["1420 Johnson St", "Smith remodel"]},
    ]
    write_fixture(
        "C17-xls-long-list",
        "Excel 97-2003 (.xls, BIFF8) long list with real date cells.",
        "input.xls",
        buf.getvalue(),
        expected("C", "input.xls", "long_list", assignments, flags,
                 {"settings": {"today": "2026-10-10", "workdays": ["mon", "tue", "wed"], "days_off": []}}),
    )


# --------------------------------------------------------------------- xlsm

def build_xlsm():
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Crew"
    ws.append(["Name"] + [datetime.fromisoformat(d) for d in DATES])
    for p, cells in PERSON_GRID.items():
        ws.append([p] + cells)
    buf = io.BytesIO()
    wb.save(buf)
    # Re-label the package as macro-enabled, as Excel does for .xlsm.
    src = zipfile.ZipFile(io.BytesIO(buf.getvalue()))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == "[Content_Types].xml":
                data = data.replace(
                    b"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml",
                    b"application/vnd.ms-excel.sheet.macroEnabled.main+xml",
                )
            z.writestr(item, data)
    write_fixture(
        "C18-xlsm-person-by-day",
        "Macro-enabled Excel (.xlsm) person × day grid; read for values only, no macros run.",
        "input.xlsm",
        out.getvalue(),
        expected("C", "input.xlsm", "person_by_day", person_assignments(PERSON_GRID), person_flags(PERSON_GRID)),
    )


# --------------------------------------------------------------------- xlsb

def rec(rtype, payload=b""):
    def varint(n, max_bytes):
        out = bytearray()
        for _ in range(max_bytes):
            b = n & 0x7F
            n >>= 7
            out.append(b | (0x80 if n else 0))
            if not n:
                break
        return bytes(out)

    return varint(rtype, 2) + varint(len(payload), 4) + payload


def wide(s):
    data = s.encode("utf-16-le")
    return struct.pack("<I", len(data) // 2) + data


def xlsb_bytes(rows, sheet_name="Schedule"):
    workbook = (
        rec(0x0083)  # BrtBeginBook
        + rec(0x0099, struct.pack("<I", 0) + struct.pack("<I", 0) + wide(""))  # BrtWbProp
        + rec(0x008F)  # BrtBeginBundleShs
        + rec(0x009C, struct.pack("<II", 0, 1) + wide("rId1") + wide(sheet_name))  # BrtBundleSh
        + rec(0x0090)  # BrtEndBundleShs
        + rec(0x0084)  # BrtEndBook
    )
    ncols = max(len(r) for r in rows)
    sheet = rec(0x0081) + rec(0x0094, struct.pack("<IIII", 0, len(rows) - 1, 0, ncols - 1)) + rec(0x0091)
    for r, row in enumerate(rows):
        sheet += rec(0x0000, struct.pack("<IIHBBBI", r, 0, 300, 0, 0, 0, 0))  # BrtRowHdr
        for c, v in enumerate(row):
            if v in ("", None):
                continue
            if isinstance(v, (int, float)):
                sheet += rec(0x0005, struct.pack("<II", c, 0) + struct.pack("<d", float(v)))  # BrtCellReal
            else:
                sheet += rec(0x0006, struct.pack("<II", c, 0) + wide(v))  # BrtCellSt
    sheet += rec(0x0092) + rec(0x0082)  # BrtEndSheetData, BrtEndSheet
    content_types = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
        '<Default Extension="bin" ContentType="application/vnd.ms-excel.sheet.binary.macroEnabled.main"/>'
        '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
        '<Override PartName="/xl/worksheets/sheet1.bin" ContentType="application/vnd.ms-excel.worksheet"/>'
        "</Types>"
    )
    root_rels = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="xl/workbook.bin"/>'
        "</Relationships>"
    )
    wb_rels = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet1.bin"/>'
        "</Relationships>"
    )
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for name, data in [
            ("[Content_Types].xml", content_types.encode()),
            ("_rels/.rels", root_rels.encode()),
            ("xl/workbook.bin", workbook),
            ("xl/_rels/workbook.bin.rels", wb_rels.encode()),
            ("xl/worksheets/sheet1.bin", sheet),
        ]:
            info = zipfile.ZipInfo(name, date_time=(2026, 10, 9, 12, 0, 0))
            z.writestr(info, data)
    return out.getvalue()


def build_xlsb():
    rows = [["Name", "Mon 10/12", "Tue 10/13", "Wed 10/14", "Thu 10/15", "Fri 10/16"]]
    rows += [[p] + cells for p, cells in PERSON_GRID.items()]
    write_fixture(
        "C19-xlsb-person-by-day",
        "Excel binary workbook (.xlsb) person × day grid with text day headers.",
        "input.xlsb",
        xlsb_bytes(rows),
        expected("C", "input.xlsb", "person_by_day", person_assignments(PERSON_GRID), person_flags(PERSON_GRID)),
    )


if __name__ == "__main__":
    build_numbers()
    build_xls()
    build_xlsm()
    build_xlsb()
    print("built C15-C19")
