"""Shared helpers for writing handmade fixtures.

These helpers are authoring tools, not a schedule reader. Each fixture's
truth is written by hand in a build_<category>.py script. The helpers render
that truth into the input file and spell out the matching `assignments`, so
the two can't drift apart. Flags and info items are always written by hand in
the build scripts and then re-derived independently by check_flags.py.

Weekday labels are produced with datetime here, never typed by hand, except
where a fixture plants a wrong weekday on purpose.
"""
from __future__ import annotations

import datetime as dt
import json
import re
import shutil
import zipfile
from pathlib import Path

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent  # fixtures/handmade
HINTS = HERE / "hints"

TODAY = "2026-10-10"
WD5 = ["mon", "tue", "wed", "thu", "fri"]
WD6 = WD5 + ["sat"]

FIXED_TIME = (2026, 10, 9, 12, 0, 0)  # zip entry timestamps, so binaries rebuild byte-identical


# ---------------------------------------------------------------- dates

def d(s: str) -> str:
    """'10-12' -> '2026-10-12'; full ISO dates pass through."""
    return s if len(s) == 10 else "2026-" + s


def days(start: str, n: int) -> list[str]:
    s = dt.date.fromisoformat(d(start))
    return [(s + dt.timedelta(i)).isoformat() for i in range(n)]


def wd(iso: str) -> str:
    return dt.date.fromisoformat(iso).strftime("%a")


def mdyyyy(iso: str) -> str:
    x = dt.date.fromisoformat(iso)
    return f"{x.month}/{x.day}/{x.year}"


def mdyy(iso: str) -> str:
    x = dt.date.fromisoformat(iso)
    return f"{x.month}/{x.day}/{x.year % 100:02d}"


def md(iso: str) -> str:
    x = dt.date.fromisoformat(iso)
    return f"{x.month}/{x.day}"


def mon_dd(iso: str) -> str:  # "Oct 12"
    x = dt.date.fromisoformat(iso)
    return f"{x.strftime('%b')} {x.day}"


def dd_mon(iso: str) -> str:  # "12-Oct" (Excel d-mmm)
    x = dt.date.fromisoformat(iso)
    return f"{x.day}-{x.strftime('%b')}"


def wd_md(iso: str) -> str:  # "Mon 10/12"
    return f"{wd(iso)} {md(iso)}"


def ddmmyyyy(iso: str) -> str:  # "13/10/2026"
    x = dt.date.fromisoformat(iso)
    return f"{x.day}/{x.month}/{x.year}"


def serial(iso: str) -> int:  # Excel 1900 date system
    return (dt.date.fromisoformat(iso) - dt.date(1899, 12, 30)).days


def pydate(iso: str) -> dt.datetime:
    x = dt.date.fromisoformat(iso)
    return dt.datetime(x.year, x.month, x.day)


# ---------------------------------------------------------------- records

def W(person, date, site, **kw):
    r = {"person": person, "date": d(date), "site": site}
    r.update(kw)
    return r


def ST(person, date, status, **kw):
    r = {"person": person, "date": d(date), "status": status}
    r.update(kw)
    return r


def UN(date, site):
    return {"date": d(date), "site": site, "unstaffed": True}


STATUS_TEXT = {
    "off": "off", "x": "off", "—": "off", "-": "off", "n/a": "off", "day off": "off",
    "pto": "pto", "vac": "vacation", "vacation": "vacation", "sick": "sick",
    "holiday": "holiday", "out": "out",
    "descanso": "off", "vacaciones": "vacation", "enfermo": "sick", "libre": "off",
}


def split_note(s: str):
    s = " ".join(s.split())
    if s.endswith(")") and "(" in s:
        i = s.rfind("(")
        return s[:i].strip(), s[i + 1:-1].strip()
    if " - " in s:
        a, b = s.split(" - ", 1)
        return a.strip(), b.strip()
    return s, None


def cell_recs(person, date, cell, site_sep=" / ", **kw):
    """Records for one person x day cell, as this fixture library writes them."""
    t = " ".join((cell or "").split())
    if not t:
        return []
    out = []
    for part in t.split(site_sep):
        if part.strip().lower() in STATUS_TEXT:
            out.append(ST(person, date, STATUS_TEXT[part.strip().lower()], **kw))
            continue
        site, note = split_note(part)
        extra = dict(kw)
        if note:
            extra["note"] = note
        out.append(W(person, date, site, **extra))
    return out


def person_grid(rows, dates, site_sep=" / "):
    """rows: [(person, [cell per date])] -> records"""
    out = []
    for person, cells in rows:
        assert len(cells) == len(dates), (person, cells)
        for date, cell in zip(dates, cells):
            out += cell_recs(person, date, cell, site_sep)
    return out


NAME_SPLIT = re.compile(r"\s*(?:,|&|/|\+|\n| and | y )\s*")


def split_names(cell: str) -> list[str]:
    return [n.strip() for n in NAME_SPLIT.split(cell or "") if n.strip()]


def job_grid(rows, dates):
    """rows: [(job label, [cell per date])]; a status label ("PTO", "Off") makes a status row."""
    out = []
    for job, cells in rows:
        assert len(cells) == len(dates), (job, cells)
        status = STATUS_TEXT.get(job.strip().lower())
        for date, cell in zip(dates, cells):
            for name in split_names(cell):
                out.append(ST(name, date, status) if status else W(name, date, job))
    return out


def grid_to_list_rows(rows, dates, fmt, site_sep=" / ", person_first=False):
    """Render a person x day truth as long-list rows (date, person, site), sorted by date."""
    out = []
    for i, date in enumerate(dates):
        for person, cells in rows:
            t = cells[i].strip()
            if not t:
                continue
            for part in t.split(site_sep):
                out.append([person, fmt(date), part] if person_first else [fmt(date), person, part])
    return out


# ---------------------------------------------------------------- flags and info (hand-written in build scripts)

def DB(person, date, *sites):
    return {"type": "double_booked", "person": person, "date": d(date), "sites": list(sites)}


def AWO(person, date, sites, status):
    return {"type": "assigned_while_off", "person": person, "date": d(date), "sites": list(sites), "status": status}


def NB(person, date):
    return {"type": "not_booked", "person": person, "date": d(date)}


def UF(site, date):
    return {"type": "unstaffed", "site": site, "date": d(date)}


def SPLIT(person, date, *sites):
    return {"type": "split_day", "person": person, "date": d(date), "sites": list(sites)}


def DUP(person, date, site):
    return {"type": "duplicate_entry", "person": person, "date": d(date), "sites": [site]}


def INTERNAL(person, date, *sites):
    return {"type": "internal_plus_site", "person": person, "date": d(date), "sites": list(sites)}


# ---------------------------------------------------------------- questions

def Q(type_, **fields):
    q = {"type": type_}
    q.update(fields)
    return q


# ---------------------------------------------------------------- file bytes

def csv_bytes(rows, delim=",", eol="\n", encoding="utf-8", bom=False, final_eol=True):
    lines = []
    for row in rows:
        cells = []
        for c in row:
            c = "" if c is None else str(c)
            if any(ch in c for ch in (delim, '"', "\n", "\r")):
                c = '"' + c.replace('"', '""') + '"'
            cells.append(c)
        lines.append(delim.join(cells))
    text = eol.join(lines) + (eol if final_eol else "")
    if encoding in ("utf-16-le", "utf-16le"):
        return (b"\xff\xfe" if bom else b"") + text.encode("utf-16-le")
    data = text.encode(encoding)
    return (b"\xef\xbb\xbf" + data) if bom else data


def settings(workdays=WD5, today=TODAY, days_off=()):
    return {"today": today, "workdays": list(workdays), "days_off": list(days_off)}


def normalize_zip(path: Path):
    """Rewrite a zip container with fixed timestamps, keeping entry order and compression.
    openpyxl stamps docProps/core.xml with the save time, so that is pinned too."""
    with zipfile.ZipFile(path) as z:
        entries = [(i, z.read(i.filename)) for i in z.infolist()]
    pinned = []
    for info, data in entries:
        if info.filename == "docProps/core.xml":
            data = re.sub(rb"(<dcterms:(?:created|modified)[^>]*>)[^<]*(</dcterms:)",
                          rb"\g<1>2026-10-09T12:00:00Z\g<2>", data)
        pinned.append((info, data))
    entries = pinned
    tmp = path.with_suffix(path.suffix + ".tmp")
    with zipfile.ZipFile(tmp, "w") as z:
        for info, data in entries:
            ni = zipfile.ZipInfo(info.filename, date_time=FIXED_TIME)
            ni.compress_type = info.compress_type
            ni.external_attr = info.external_attr
            z.writestr(ni, data)
    tmp.replace(path)


# ---------------------------------------------------------------- writing

KEY_ORDER = [
    "category", "tier", "input", "paste", "settings", "layout", "arrangement",
    "assignments", "assignments_exact", "flags", "flags_exact", "info", "info_exact",
    "ask", "must_not_ask", "answers", "after_answers", "error", "notes",
]


def expected(category, input, *, tier="script", paste=False, settings_=None, layout=None,
             arrangement=None, assignments=None, assignments_exact=True, flags=None,
             info=None, ask=None, must_not_ask=None, answers=None, after_answers=None,
             notes=None):
    e = {
        "category": category,
        "tier": tier,
        "input": input,
        "paste": paste,
        "settings": settings_ if settings_ is not None else settings(),
        "layout": layout,
        "arrangement": arrangement,
        "assignments": assignments,
        "assignments_exact": assignments_exact if assignments is not None else None,
        "flags": flags,
        "flags_exact": True if flags is not None else None,
        "info": info,
        "ask": ask,
        "must_not_ask": must_not_ask,
        "answers": answers,
        "after_answers": after_answers,
        "notes": notes,
    }
    return {k: e[k] for k in KEY_ORDER if e.get(k) is not None}


def dumps(v, indent=0) -> str:
    """JSON with objects expanded and lists of records written one record per line."""
    pad = "  " * indent
    if isinstance(v, dict) and indent < 2 and v and any(isinstance(x, (dict, list)) for x in v.values()):
        parts = [f'{pad}  {json.dumps(k, ensure_ascii=False)}: {dumps(x, indent + 1).lstrip()}' for k, x in v.items()]
        return pad + "{\n" + ",\n".join(parts) + "\n" + pad + "}"
    if isinstance(v, list) and v and all(isinstance(x, dict) for x in v):
        parts = [f'{pad}  {json.dumps(x, ensure_ascii=False)}' for x in v]
        return pad + "[\n" + ",\n".join(parts) + "\n" + pad + "]"
    return pad + json.dumps(v, ensure_ascii=False)


def emit(fid: str, slug: str, readme: str, files: dict, exp: dict, hints: dict | None = None):
    """files: {name: bytes | callable(path)}; a callable writes the file itself (binaries)."""
    assert re.fullmatch(r"[A-N]\d\d", fid), fid
    assert re.fullmatch(r"[a-z0-9]+(?:-[a-z0-9]+)*", slug), slug
    assert "\n" not in readme.strip()
    for old in ROOT.glob(f"{fid}-*"):
        shutil.rmtree(old)
    folder = ROOT / f"{fid}-{slug}"
    folder.mkdir(parents=True)
    (folder / "README.md").write_text(readme.strip() + "\n", encoding="utf-8")
    for name, content in files.items():
        p = folder / name
        if callable(content):
            content(p)
            if p.suffix in (".xlsx", ".ods", ".docx"):
                normalize_zip(p)
        else:
            p.write_bytes(content)
    assert exp["input"] in files, (fid, exp["input"])
    text = dumps(exp) + "\n"
    assert json.loads(text) == exp, fid
    (folder / "expected.json").write_text(text, encoding="utf-8")
    HINTS.mkdir(exist_ok=True)
    hp = HINTS / f"{fid}.json"
    if hints:
        hp.write_text(json.dumps(hints, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    elif hp.exists():
        hp.unlink()
    return folder


# ---------------------------------------------------------------- G, H, J, K helpers
# Added for the status / times / scale / flag-logic fixtures. Append-only: nothing
# above changes, so the A-D, I and M builds stay byte-identical.

UNSTAFFED_TEXT = {"tbd", "open", "?", "need guy", "need 1"}
INACTIVE_TEXT = {"x", "—", "n/a", "no work", "closed", "done"}  # job-grid cells only
PART_TEXT = {"am": "am", "am only": "am", "pm": "pm", "pm only": "pm", "half day": "half", "half": "half"}

_CLOCK = re.compile(r"^(\d{1,2})(?::(\d{2}))?\s*([ap])?m?$", re.I)
_RANGE = re.compile(r"^(.*?)\s*[-–]\s*(.*)$")


def _clock(text):
    m = _CLOCK.match(text.strip())
    assert m, f"not a clock time: {text!r}"
    h, mm, ap = int(m.group(1)), int(m.group(2) or 0), (m.group(3) or "").lower()
    return h, mm, ap


def time_range(text: str):
    """'7-11' -> ('07:00','11:00'); '7a-3p'; '7:00–15:30'; '10p-6a'; '12-4' -> 12:00-16:00.
    Hours without am/pm follow DECISIONS.md: a start of 5-11 is morning, 12 is noon,
    1-4 is afternoon; the end is the first matching hour after the start."""
    m = _RANGE.match(text.strip())
    assert m, text
    (h1, m1, a1), (h2, m2, a2) = _clock(m.group(1)), _clock(m.group(2))
    if a1:
        s = (h1 % 12 + (12 if a1 == "p" else 0)) * 60 + m1
    elif h1 >= 13:
        s = h1 * 60 + m1
    else:
        s = (h1 + (12 if 1 <= h1 <= 4 else 0)) * 60 + m1
    if a2:
        e = (h2 % 12 + (12 if a2 == "p" else 0)) * 60 + m2
    elif h2 >= 13:
        e = h2 * 60 + m2
    else:
        cands = [(h2 % 12) * 60 + m2, (h2 % 12 + 12) * 60 + m2]
        later = [c for c in cands if c > s]
        e = min(later) if later else cands[0]
    return f"{s // 60:02d}:{s % 60:02d}", f"{e // 60:02d}:{e % 60:02d}"


_TRAIL_TIME = re.compile(r"^(.*\S)\s+(\d{1,2}(?::\d{2})?\s*(?:[ap]m?)?\s*[-–]\s*\d{1,2}(?::\d{2})?\s*(?:[ap]m?)?)$", re.I)


def entry_recs(person, date, entry, **kw):
    """One entry of a person-grid cell (cells split on ' / ' first):
    status word, optionally with a trailing note ('PTO (dentist)', 'sick - flu');
    or a site with an optional trailing part ('(AM only)', ' - half day'),
    trailing time range ('Smith remodel 7-11') and trailing note."""
    t = " ".join(entry.split())
    if not t:
        return []
    base, note = split_note(t)
    if base.lower() in STATUS_TEXT:
        r = ST(person, date, STATUS_TEXT[base.lower()], **kw)
        if note:
            r["note"] = note
        return [r]
    extra = dict(kw)
    if note is not None and note.lower() in PART_TEXT:
        extra["part"] = PART_TEXT[note.lower()]
        base, note = split_note(base)
    m = _TRAIL_TIME.match(base)
    if m:
        base = m.group(1)
        extra["start"], extra["end"] = time_range(m.group(2))
    assert base.lower() not in UNSTAFFED_TEXT and base.lower() not in PART_TEXT, t
    if note:
        extra["note"] = note
    return [W(person, date, base, **extra)]


def person_cells(rows, dates, sep=" / "):
    """rows: [(person, [cell per date])] -> records, using entry_recs per entry."""
    out = []
    for person, cells in rows:
        assert len(cells) == len(dates), (person, cells)
        for date, cell in zip(dates, cells):
            for entry in (cell or "").split(sep):
                out += entry_recs(person, date, entry)
    return out


def job_cells(rows, dates):
    """Job x day grid -> (records, inactive [(job, date)]).
    Unstaffed markers give an unstaffed record; inactive markers give nothing;
    a status label row ('PTO', 'Off') gives status records."""
    out, inactive = [], []
    for job, cells in rows:
        assert len(cells) == len(dates), (job, cells)
        status = STATUS_TEXT.get(job.strip().lower())
        for date, cell in zip(dates, cells):
            t = " ".join((cell or "").split())
            if not t:
                continue
            if not status and t.lower() in UNSTAFFED_TEXT:
                out.append(UN(date, job))
                continue
            if not status and t.lower() in INACTIVE_TEXT:
                inactive.append([job, date])
                continue
            for name in split_names(t):
                assert name.lower() not in UNSTAFFED_TEXT | INACTIVE_TEXT, (job, cell)
                out.append(ST(name, date, status) if status else W(name, date, job))
    return out, inactive


def SUBMULTI(person, date, *sites):
    return {"type": "subcontractor_multi_site", "person": person, "date": d(date), "sites": list(sites)}


def xlsx_workbook(creator="Office"):
    import openpyxl
    wb = openpyxl.Workbook()
    wb.properties.creator = creator
    wb.properties.created = wb.properties.modified = dt.datetime(*FIXED_TIME)
    return wb


# ---------------------------------------------------------------- E, F helpers
# Added for the name and site matching fixtures. Append-only, like the block above.

def ANS(question, answer):
    """One entry of `answers`."""
    return {"question": question, "answer": answer}


def plan_rows(plan, dates):
    """plan: [(person, [cell per date])]; a cell is '', a site, or a list of sites.
    -> [(iso date, person, site)] in file order: by date, then plan order."""
    out = []
    for i, x in enumerate(dates):
        for person, cells in plan:
            assert len(cells) == len(dates), (person, cells)
            c = cells[i]
            for s in ([c] if isinstance(c, str) else c):
                if s:
                    out.append((x, person, s))
    return out


def list_file(rows, header=("Date", "Name", "Job")):
    """rows: [(iso date, person as typed, site as typed)] -> (lines, assignments).
    Assignments carry the trimmed, whitespace-collapsed text (FORMAT.md)."""
    lines = [list(header)] + [[mdyyyy(x), p, s] for x, p, s in rows]
    return lines, [W(" ".join(p.split()), x, " ".join(s.split())) for x, p, s in rows]


def day_lines(cols, dates, first="Date"):
    """Transposed grid: dates down the side, column labels (jobs or people) across."""
    return [[first] + [label for label, _ in cols]] + [[mdyyyy(x)] + [cells[i] for _, cells in cols]
                                                       for i, x in enumerate(dates)]


def write_lines_xlsx(path, lines, date_axis):
    """Write csv-style lines to a one-sheet workbook. date_axis 'row': the header
    row's cells after the first are real dates; 'col': the first column's cells
    below the header are real dates. Blank strings stay empty cells."""
    wb = xlsx_workbook()
    ws = wb.active
    ws.title = "Schedule"
    for r, line in enumerate(lines, start=1):
        for c, v in enumerate(line, start=1):
            if v in ("", None):
                continue
            if (date_axis == "row" and r == 1 and c > 1) or (date_axis == "col" and r > 1 and c == 1):
                m, dd, y = (int(t) for t in v.split("/"))
                ws.cell(row=r, column=c, value=pydate(f"{y:04d}-{m:02d}-{dd:02d}")).number_format = "m/d/yyyy"
            else:
                ws.cell(row=r, column=c, value=v)
    wb.save(path)
