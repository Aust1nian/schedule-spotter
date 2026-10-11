"""Category B: headers. Run: uv run --with openpyxl python3 build_B.py"""
import datetime as dt

from fxlib import *

WEEK = days("10-12", 5)
J, S, O = "1420 Johnson St", "Smith remodel", "Oak Ridge Apts"
NO_LAYOUT_QS = [Q("pick_layout"), Q("column_role")]
SPANISH_DAYS = ["Lunes", "Martes", "Miércoles", "Jueves", "Viernes", "Sábado", "Domingo"]


def b01():
    rows = [
        ("Mike Baker", [J] * 5),
        ("Jose Garcia", [S, S, S, f"{S} / off", S]),
        ("Luis Ortiz", [O] * 5),
        ("Ana Ruiz", [O, O, S, S, S]),
    ]
    body = []
    for person, cells in rows:  # grouped by person, then date
        for date, cell in zip(WEEK, cells):
            for part in cell.split(" / "):
                body.append([person, mdyyyy(date), part])
    exp = expected("B", "input.csv", layout="long_list", arrangement="single",
                   assignments=person_grid(rows, WEEK),
                   flags=[AWO("Jose Garcia", "10-15", [S], "off")], info=[],
                   must_not_ask=NO_LAYOUT_QS)
    emit("B01", "no-header", "No header row at all: columns (name, date, job) must be told apart by their values alone.",
         {"input.csv": csv_bytes(body)}, exp)


def b02():
    rows = [
        ("Mike Baker", [J] * 5),
        ("Jose Garcia", [S, S, f"{S} / {O}", S, S]),
        ("Luis Ortiz", [O] * 5),
        ("Dave Kim", [J] * 5),
        ("Tom Reed", [S, S, S, S, "off"]),
    ]
    blank = [""] * 6
    lines = [
        ["Ridgeline Builders"] + [""] * 5,
        ["Crew Schedule"] + [""] * 5,
        ["Week of Oct 12, 2026"] + [""] * 5,
        ["Printed 10/9/2026"] + [""] * 5,
        blank,
        ["Name"] + [mdyyyy(x) for x in WEEK],
    ] + [[p] + c for p, c in rows]
    exp = expected("B", "input.csv", layout="person_by_day", arrangement="single",
                   assignments=person_grid(rows, WEEK),
                   flags=[DB("Jose Garcia", "10-14", S, O)], info=[],
                   must_not_ask=NO_LAYOUT_QS)
    emit("B02", "title-rows-above-header",
         "Four title rows (company, sheet title, 'Week of', a 'Printed 10/9/2026' date) and a blank row above the real header; none are data and 10/9 must not join the window.",
         {"input.csv": csv_bytes(lines, eol="\r\n")}, exp, {"window": WEEK})


def b03():
    import openpyxl
    from openpyxl.styles import Alignment

    halves = {  # AM, PM per day
        "Mike Baker": [J, J, J, J, J, S, S, S, S, S],
        "Jose Garcia": [S] * 10,
        "Luis Ortiz": [O] * 10,
        "Dave Kim": [O] * 10,
    }
    assigns = []
    for person, cells in halves.items():
        for i, site in enumerate(cells):
            if site:
                assigns.append(W(person, WEEK[i // 2], site, part="am" if i % 2 == 0 else "pm"))

    def make(path):
        wb = openpyxl.Workbook()
        wb.properties.creator = "Ridgeline Builders"
        wb.properties.created = wb.properties.modified = dt.datetime(2026, 10, 9, 12, 0)
        ws = wb.active
        ws.title = "Schedule"
        ws["A1"] = "Name"
        ws.merge_cells("A1:A2")
        for i, date in enumerate(WEEK):
            col = 2 + 2 * i
            c = ws.cell(row=1, column=col, value=pydate(date))
            c.number_format = "dddd m/d"
            c.alignment = Alignment(horizontal="center")
            ws.merge_cells(start_row=1, start_column=col, end_row=1, end_column=col + 1)
            ws.cell(row=2, column=col, value="AM")
            ws.cell(row=2, column=col + 1, value="PM")
        for r, (person, cells) in enumerate(halves.items(), start=3):
            ws.cell(row=r, column=1, value=person)
            for i, site in enumerate(cells):
                if site:
                    ws.cell(row=r, column=2 + i, value=site)
        wb.save(path)

    exp = expected("B", "input.xlsx", layout="person_by_day", arrangement="single",
                   assignments=assigns, flags=[],
                   info=[SPLIT("Mike Baker", "10-14", J, S)],
                   must_not_ask=NO_LAYOUT_QS + [Q("pick_sheet")])
    emit("B03", "two-row-header-merged-am-pm",
         "Two-row header from merged cells in .xlsx: each date (a real date cell shown as 'Monday 10/12') is merged across AM and PM sub-columns.",
         {"input.xlsx": make}, exp, {"window": WEEK})


def b04():
    plan = {
        "10-12": {J: ["Mike Baker", "Dave Kim"], S: ["Jose Garcia", "Luis Ortiz", "Ana Ruiz"], O: ["Tom Reed"]},
        "10-13": {J: ["Mike Baker", "Dave Kim"], S: ["Jose Garcia", "Luis Ortiz", "Ana Ruiz"], O: ["Tom Reed", "Luis Ortiz"]},
        "10-14": {J: ["Mike Baker", "Dave Kim"], S: ["Jose Garcia", "Ana Ruiz"], O: ["Tom Reed", "Luis Ortiz"]},
        "10-15": {J: ["Mike Baker"], S: ["Jose Garcia", "Ana Ruiz", "Dave Kim"], O: ["Tom Reed", "Luis Ortiz"]},
        "10-16": {J: ["Mike Baker", "Dave Kim"], S: ["Jose Garcia", "Ana Ruiz"], O: ["Tom Reed", "Luis Ortiz"]},
    }
    lines = [["Date", "Job", "Name", "Name", "Name"]]
    assigns = []
    for day, jobs in plan.items():
        for job, names in jobs.items():
            lines.append([mdyyyy(d(day)), job] + names + [""] * (3 - len(names)))
            assigns += [W(n, day, job) for n in names]
    exp = expected("B", "input.csv", layout="long_list", arrangement="single",
                   assignments=assigns, flags=[DB("Luis Ortiz", "10-13", S, O)], info=[],
                   must_not_ask=NO_LAYOUT_QS + [Q("fill_down")])
    emit("B04", "duplicate-column-names",
         "Duplicate column names: three columns all headed 'Name' (crew slots on one job-day row), each a person column.",
         {"input.csv": csv_bytes(lines)}, exp)


def b05():
    P, E, L = "Parker kitchen", "88 Elm Ave", "Lincoln School"
    rows = [
        ("Ana Ruiz", [P, f"{P} / sick", P, P, P]),
        ("Carlos Mendez", [E] * 5),
        ("Pedro Sanchez", [L] * 5),
        ("Kevin Walsh", [P, P, P, E, E]),
    ]
    header = ["Date", "Name", "Job"]
    body = grid_to_list_rows(rows, WEEK, mdyyyy)
    lines = []
    for i, r in enumerate(body):
        if i % 8 == 0:  # header repeats at every printed page break
            lines.append(header)
        lines.append(r)
    exp = expected("B", "input.csv", layout="long_list", arrangement="single",
                   assignments=person_grid(rows, WEEK),
                   flags=[AWO("Ana Ruiz", "10-13", [P], "sick")], info=[],
                   must_not_ask=NO_LAYOUT_QS)
    emit("B05", "header-repeated-mid-file",
         "The header row repeats every 8 data rows (copied from a printed report); repeated headers are not data.",
         {"input.csv": csv_bytes(lines)}, exp)


def b06():
    H, LV, M = "Henderson", "Lakeview HOA", "Murphy"
    rows = [
        ("Ray Jones", [H] * 5),
        ("Sam Patel", [LV, LV, LV, M, M]),
        ("Eddie Flores", [M] * 5),
    ]
    lines = [["Day", "Tech", "Customer"]] + grid_to_list_rows(rows, WEEK, mdyyyy)
    exp = expected("B", "input.csv", layout="long_list", arrangement="single",
                   assignments=person_grid(rows, WEEK), flags=[], info=[],
                   must_not_ask=NO_LAYOUT_QS)
    emit("B06", "synonyms-tech-customer",
         "Header synonyms 'Tech' (person) and 'Customer' (site), service-company style, with 'Day' holding full dates.",
         {"input.csv": csv_bytes(lines)}, exp)


def b07():
    R, L = "Riverside duplex", "Lincoln School"
    rows = [
        ("Juan Lopez", [O] * 5),
        ("Miguel Torres", [R, R, R, f"{R} / {L}", R]),
        ("Chris Hall", [L] * 5),
        ("Brian Young", [O, O, O, R, R]),
    ]
    lines = [["Project", "Date", "Worker"]]
    for project in (O, R, L):  # grouped by project, then date
        for i, date in enumerate(WEEK):
            for person, cells in rows:
                if project in cells[i].split(" / "):
                    lines.append([project, mdyyyy(date), person])
    exp = expected("B", "input.csv", layout="long_list", arrangement="single",
                   assignments=person_grid(rows, WEEK),
                   flags=[DB("Miguel Torres", "10-15", R, L)], info=[],
                   must_not_ask=NO_LAYOUT_QS)
    emit("B07", "synonyms-worker-project",
         "Header synonyms 'Worker' (person) and 'Project' (site), with the site column first.",
         {"input.csv": csv_bytes(lines)}, exp)


def b08():
    M, E = "305 Main St", "88 Elm Ave"
    rows = [
        ("Mike Baker", [J] * 5),
        ("Dave Kim", [M] * 5),
        ("Tom Reed", [M, M, E, E, E]),
        ("Luis Ortiz", [E, E, E, E, ""]),
    ]
    lines = [["Date", "Crew", "Address"]] + grid_to_list_rows(rows, WEEK, mdyyyy)
    exp = expected("B", "input.csv", layout="long_list", arrangement="single",
                   assignments=person_grid(rows, WEEK),
                   flags=[NB("Luis Ortiz", "10-16")], info=[],
                   must_not_ask=NO_LAYOUT_QS)
    emit("B08", "synonyms-crew-address",
         "Header synonyms 'Crew' (a person column holding individual names, not crew names) and 'Address' (site).",
         {"input.csv": csv_bytes(lines)}, exp)


def b09():
    CJ, RS = "Casa Johnson", "Remodelación Smith"
    rows = [
        ("José García", [CJ] * 5),
        ("Luis Ortiz", [RS, RS, RS, RS, "descanso"]),
        ("Pedro Sánchez", [O, O, O, f"vacaciones / {O}", "vacaciones"]),
        ("Juan López", [RS, RS, O, O, O]),
    ]
    lines = [["Fecha", "Empleado", "Obra"]] + grid_to_list_rows(rows, WEEK, mdyyyy)
    exp = expected("B", "input.csv", layout="long_list", arrangement="single",
                   assignments=person_grid(rows, WEEK),
                   flags=[AWO("Pedro Sánchez", "10-15", [O], "vacation")], info=[],
                   must_not_ask=NO_LAYOUT_QS)
    emit("B09", "spanish-headers-long-list",
         "Spanish headers Fecha / Empleado / Obra with Spanish status words (descanso, vacaciones); US month-first dates.",
         {"input.csv": csv_bytes(lines)}, exp)


def b10():
    rows = [
        ("José García", ["Johnson"] * 5),
        ("Luis Ortiz", ["Smith"] * 5),
        ("Pedro Sánchez", ["Oak Ridge"] * 4 + ["libre"]),
        ("Juan López", ["Oak Ridge", "Oak Ridge", "enfermo", "Johnson", "Johnson"]),
    ]
    header = ["Empleado"] + [f"{SPANISH_DAYS[dt.date.fromisoformat(x).weekday()]} {md(x)}" for x in WEEK]
    assert header[1] == "Lunes 10/12" and header[5] == "Viernes 10/16"
    lines = [header] + [[p] + c for p, c in rows]
    exp = expected("B", "input.csv", layout="person_by_day", arrangement="single",
                   assignments=person_grid(rows, WEEK), flags=[], info=[],
                   must_not_ask=NO_LAYOUT_QS + [Q("weekday_mismatch")])
    emit("B10", "spanish-headers-person-by-day",
         "Spanish person x day header: Empleado, then Lunes 10/12 ... Viernes 10/16 (Spanish weekday names plus M/D); statuses libre and enfermo.",
         {"input.csv": csv_bytes(lines)}, exp, {"window": WEEK})


if __name__ == "__main__":
    for f in (b01, b02, b03, b04, b05, b06, b07, b08, b09, b10):
        f()
    print("B done")
