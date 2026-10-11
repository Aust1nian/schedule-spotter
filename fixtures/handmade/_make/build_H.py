"""Category H: times.
Run: uv run --with openpyxl==3.1.5 python3 build_H.py

Times in assignments are 24-hour HH:MM; hours written without am/pm follow
DECISIONS.md ("Hours without am/pm"): a start of 5-11 is morning, 12 is noon,
1-4 is afternoon, and the end is the first matching hour after the start.

Readings chosen where FORMAT.md is silent:
  - Times are half-open: 7a-12p then 12p-4p touch at noon and do not overlap
    (H03), so that day is a split day.
  - double_booked compares records that share a date. H05's next-morning job
    starts at 7a, after the 10p-6a shift ends, so a reading that also compares
    across midnight gives the same answer.
  - A status record from a PM sub-column carries part "pm" (H02), so it does
    not conflict with that morning's work.
"""
import datetime as dt

from fxlib import *

WEEK = days("10-12", 5)
J, S, O, P, L = "1420 Johnson St", "Smith remodel", "Oak Ridge Apts", "Parker kitchen", "Lincoln School"
HWY = "Hwy 99 overpass"
NO_ASK = [Q("pick_layout"), Q("column_role"), Q("cell_meaning")]


def exact_info(exp):
    exp = dict(exp)
    exp["info_exact"] = True
    return {k: exp[k] for k in KEY_ORDER if k in exp}


def timed_list(plan, header):
    """plan: [(person, [entries per weekday])], an entry is [(site, time text or "")].
    Rows are written date by date in plan order. Returns (lines, assignments)."""
    lines, assigns = [list(header)], []
    for i, x in enumerate(WEEK):
        for person, per_day in plan:
            for site, t in per_day[i]:
                lines.append([mdyyyy(x), person, site, t])
                kw = {}
                if t:
                    kw["start"], kw["end"] = time_range(t)
                assigns.append(W(person, x, site, **kw))
    return lines, assigns


def same(site, t, n=5):
    return [[(site, t)] for _ in range(n)]


def h01():
    plan = [
        ("Mike Baker", [[(J, t)] for t in ["7-3:30", "7a-3p", "7:00–15:30", "7-11", "6:30a-3p"]]),
        ("Jose Garcia", [[(S, t)] for t in ["7:00–15:30", "7-11", "7a-3p", "6:30a-3p", "7-3:30"]]),
        ("Luis Ortiz", [[(O, t)] for t in ["7a-3p", "7:00–15:30", "7-11", "7a-3p", "7:00–15:30"]]),
        ("Dave Kim", [[(P, t)] for t in ["7-11", "6:30a-3p", "7-3:30", "7:00–15:30", "7a-3p"]]),
    ]
    lines, assigns = timed_list(plan, ["Date", "Name", "Job", "Time"])
    exp = exact_info(expected("H", "input.csv", layout="long_list", arrangement="single",
                              assignments=assigns, flags=[], info=[], must_not_ask=NO_ASK))
    emit("H01", "time-range-formats",
         "Long list with a Time column in mixed formats ('7-11', '7a-3p', '7:00–15:30' with an en dash, '7-3:30', '6:30a-3p'), each read as a 24-hour start and end.",
         {"input.csv": csv_bytes(lines)}, exp)


def h02():
    halves = {  # (AM, PM) per day
        "Mike Baker": [(J, J), (J, S), (J, J), (J, J), (J, J)],
        "Jose Garcia": [(S, S), (S, S), (S, "off"), (S, S), (S, S)],
        "Luis Ortiz": [(O, O), (O, O), (O, O), (O, ""), (O, O)],
        "Dave Kim": [(O, O), (O, O), (O, O), (O, O), (S, O)],
    }
    assigns = []
    for person, cells in halves.items():
        for x, (am, pm) in zip(WEEK, cells):
            for part, text in (("am", am), ("pm", pm)):
                assigns += entry_recs(person, x, text, part=part)
    assert {"person": "Jose Garcia", "date": "2026-10-14", "status": "off", "part": "pm"} in assigns

    def make(path):
        from openpyxl.styles import Alignment
        wb = xlsx_workbook()
        ws = wb.active
        ws.title = "Schedule"
        ws["A1"] = "Name"
        ws.merge_cells("A1:A2")
        for i, x in enumerate(WEEK):
            col = 2 + 2 * i
            c = ws.cell(row=1, column=col, value=pydate(x))
            c.number_format = "ddd m/d"
            c.alignment = Alignment(horizontal="center")
            ws.merge_cells(start_row=1, start_column=col, end_row=1, end_column=col + 1)
            ws.cell(row=2, column=col, value="AM")
            ws.cell(row=2, column=col + 1, value="PM")
        for r, (person, cells) in enumerate(halves.items(), start=3):
            ws.cell(row=r, column=1, value=person)
            for i, (am, pm) in enumerate(cells):
                if am:
                    ws.cell(row=r, column=2 + 2 * i, value=am)
                if pm:
                    ws.cell(row=r, column=3 + 2 * i, value=pm)
        wb.save(path)

    exp = exact_info(expected("H", "input.xlsx", layout="person_by_day", arrangement="single",
                              assignments=assigns, flags=[],
                              info=[SPLIT("Mike Baker", "10-13", J, S), SPLIT("Dave Kim", "10-16", S, O)],
                              must_not_ask=NO_ASK + [Q("pick_sheet")]))
    emit("H02", "am-pm-sub-columns",
         "Person x day grid (.xlsx) with each date merged over AM and PM sub-columns: different sites AM and PM are a split day, the same site twice is a full day, and 'off' in a PM cell does not conflict with that morning's work.",
         {"input.xlsx": make}, exp, {"window": WEEK})


def h03():
    plan = [
        ("Mike Baker", [[(J, "7-3:30")], [(J, "7-11"), (S, "12-3:30")], [(J, "7-3:30")], [(J, "7-3:30")], [(J, "7-3:30")]]),
        ("Jose Garcia", [[(S, "7-3:30")], [(S, "7-3:30")], [(S, "7-3:30")], [(S, "7a-12p"), (O, "12p-4p")], [(S, "7-3:30")]]),
        ("Luis Ortiz", same(O, "7-3:30")),
        ("Dave Kim", same(J, "7-3:30")),
    ]
    lines, assigns = timed_list(plan, ["Date", "Name", "Job", "Hours"])
    exp = exact_info(expected("H", "input.csv", layout="long_list", arrangement="single",
                              assignments=assigns, flags=[],
                              info=[SPLIT("Mike Baker", "10-13", J, S), SPLIT("Jose Garcia", "10-15", S, O)],
                              must_not_ask=NO_ASK))
    emit("H03", "split-day-no-overlap",
         "Two jobs in one day at times that don't overlap ('7-11' then '12-3:30'; '7a-12p' then '12p-4p', which only touch at noon) are a split day: info only, not double-booked.",
         {"input.csv": csv_bytes(lines)}, exp)


def h04():
    plan = [
        ("Mike Baker", [[(J, "7-3:30")], [(J, "7-12"), (S, "11-3")], [(J, "7-3:30")], [(J, "7-3:30")], [(J, "7-3:30")]]),
        ("Jose Garcia", same(S, "7-3:30")),
        ("Luis Ortiz", [[(O, "7-3:30")], [(O, "7-3:30")], [(O, "7-11"), (P, "11:30-3:30")], [(O, "7-3:30")], [(O, "7-3:30")]]),
        ("Dave Kim", same(P, "7-3:30")),
    ]
    lines, assigns = timed_list(plan, ["Date", "Name", "Job", "Time"])
    exp = exact_info(expected("H", "input.csv", layout="long_list", arrangement="single",
                              assignments=assigns, flags=[DB("Mike Baker", "10-13", J, S)],
                              info=[SPLIT("Luis Ortiz", "10-14", O, P)], must_not_ask=NO_ASK))
    emit("H04", "overlapping-times-flag",
         "Two jobs in one day whose times overlap (7-12 and 11-3) are double-booked; a pair that doesn't overlap (7-11 and 11:30-3:30) is only a split day.",
         {"input.csv": csv_bytes(lines)}, exp)


def h05():
    plan = [
        ("Mike Baker", same(J, "7a-3p")),
        ("Ray Jones", [[(J, "7a-3p"), (HWY, "10p-6a")]] + same(J, "7a-3p", 4)),
        ("Carlos Mendez", same(L, "2p-10p", 3) + [[(L, "2p-10p"), (HWY, "8p-4a")]] + same(L, "2p-10p", 1)),
        ("Tom Reed", same(HWY, "10p-6a")),
    ]
    lines, assigns = timed_list(plan, ["Date", "Name", "Job", "Shift"])
    assert {"person": "Ray Jones", "date": "2026-10-12", "site": HWY, "start": "22:00", "end": "06:00"} in assigns
    exp = exact_info(expected("H", "input.csv", layout="long_list", arrangement="single",
                              assignments=assigns, flags=[DB("Carlos Mendez", "10-15", L, HWY)],
                              info=[SPLIT("Ray Jones", "10-12", J, HWY)], must_not_ask=NO_ASK))
    emit("H05", "overnight-shift",
         "Overnight shifts end the next morning: 7a-3p plus that night's 10p-6a is a split day, the next day's 7a start is no conflict, and 2p-10p against 8p-4a overlaps from 8 to 10 pm, so it's double-booked.",
         {"input.csv": csv_bytes(lines)}, exp)


def h06():
    plan = [
        ("Mike Baker", [[(J, "")], [(J, "7-11"), (S, "")], [(J, "")], [(J, "")], [(J, "")]]),
        ("Jose Garcia", [[(S, "7-3:30")], [(S, "7-3:30")], [(S, "7-11"), (O, "12-4")], [(S, "")], [(S, "")]]),
        ("Luis Ortiz", same(O, "")),
        ("Dave Kim", same(J, "7-3:30")),
    ]
    lines, assigns = timed_list(plan, ["Date", "Name", "Job", "Time"])
    exp = exact_info(expected("H", "input.tsv", layout="long_list", arrangement="single",
                              assignments=assigns, flags=[DB("Mike Baker", "10-13", J, S)],
                              info=[SPLIT("Jose Garcia", "10-14", S, O)], must_not_ask=NO_ASK))
    emit("H06", "times-on-some-rows",
         "Times on only some rows (.tsv): one job with a time and another with none on the same day is double-booked; two timed jobs that don't overlap are only a split day.",
         {"input.tsv": csv_bytes(lines, delim="\t")}, exp)


def h07():
    t = lambda h, m=0: dt.time(h, m)  # noqa: E731
    plan = [  # person, [[(site, start, end)] per weekday]
        ("Mike Baker", [[(J, t(7), t(15, 30))]] * 2 + [[(J, t(7), t(11)), (S, t(11, 30), t(15, 30))]] + [[(J, t(7), t(15, 30))]] * 2),
        ("Jose Garcia", [[(S, t(7), t(15, 30))]] * 5),
        ("Luis Ortiz", [[(O, t(6, 30), t(15))]] * 5),
        ("Dave Kim", [[(P, t(7), t(15, 30))]] * 3 + [[(P, t(7), t(12)), (O, t(11, 30), t(15, 30))]] + [[(P, t(7), t(15, 30))]]),
    ]
    rows, assigns = [], []
    for i, x in enumerate(WEEK):
        for person, per_day in plan:
            for site, a, b in per_day[i]:
                rows.append((x, person, site, a, b))
                assigns.append(W(person, x, site, start=a.strftime("%H:%M"), end=b.strftime("%H:%M")))

    def make(path):
        wb = xlsx_workbook()
        ws = wb.active
        ws.title = "Schedule"
        ws.append(["Date", "Name", "Job", "Start", "End"])
        for x, person, site, a, b in rows:
            ws.append([pydate(x), person, site, a, b])
            r = ws.max_row
            ws.cell(row=r, column=1).number_format = "m/d/yyyy"
            ws.cell(row=r, column=4).number_format = "h:mm AM/PM"
            ws.cell(row=r, column=5).number_format = "h:mm AM/PM"
        wb.save(path)

    exp = exact_info(expected("H", "input.xlsx", layout="long_list", arrangement="single",
                              assignments=assigns, flags=[DB("Dave Kim", "10-15", P, O)],
                              info=[SPLIT("Mike Baker", "10-14", J, S)],
                              must_not_ask=NO_ASK + [Q("pick_sheet")]))
    emit("H07", "excel-time-cells-start-end",
         "Long list (.xlsx) with Start and End columns holding real Excel time cells shown as '7:00 AM' / '3:30 PM': read as 24-hour times, so 7:00 AM-12:00 PM against 11:30 AM-3:30 PM is double-booked.",
         {"input.xlsx": make}, exp)


def h08():
    rows = [
        ("Mike Baker", [f"{J} 7-3:30"] * 5),
        ("Jose Garcia", [f"{S} 7-3:30", f"{S} 7-11 / {P} 12-4", f"{S} 7-3:30", f"{S} 7-3:30", f"{S} 7-3:30"]),
        ("Luis Ortiz", [f"{O} 6-2:30"] * 3 + [f"{O} 6-12 / {P} 11-3", f"{O} 6-2:30"]),
        ("Dave Kim", [f"{P} 7-3:30"] * 5),
    ]
    assigns = person_cells(rows, WEEK)
    assert all("start" in a for a in assigns)
    lines = [["Name"] + [mdyyyy(x) for x in WEEK]] + [[p] + c for p, c in rows]
    exp = exact_info(expected("H", "input.csv", layout="person_by_day", arrangement="single",
                              assignments=assigns, flags=[DB("Luis Ortiz", "10-15", O, P)],
                              info=[SPLIT("Jose Garcia", "10-13", S, P)], must_not_ask=NO_ASK))
    emit("H08", "times-inside-grid-cells",
         "Person x day grid with the time written after the site in each cell ('Smith remodel 7-11 / Parker kitchen 12-4'): the time is split off the site name; overlapping times are double-booked, others a split day.",
         {"input.csv": csv_bytes(lines)}, exp, {"window": WEEK})


if __name__ == "__main__":
    for f in (h01, h02, h03, h04, h05, h06, h07, h08):
        f()
    print("H done")
