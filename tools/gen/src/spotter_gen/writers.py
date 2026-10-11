"""File writers. Every writer is byte-for-byte deterministic: zip-based formats are re-zipped
with fixed timestamps and fixed document-property dates."""

from __future__ import annotations

import csv
import datetime as dt
import io
import re
import zipfile

from .plan import Plan
from .sheet import Band, Canvas

FIXED_ZIP_TIME = (2026, 10, 9, 8, 0, 0)
FIXED_ISO = "2026-10-09T08:00:00Z"
FIXED_DT = dt.datetime(2026, 10, 9, 8, 0, 0)


def cell_str(v) -> str:
    if v is None:
        return ""
    if isinstance(v, bool):
        return str(v)
    if isinstance(v, int):
        return str(v)
    if isinstance(v, float):
        return repr(v)
    if isinstance(v, dt.date):
        return f"{v.month}/{v.day}/{v.year}"
    return str(v)


def canvas_rows(cv: Canvas, pad: bool) -> list[list[str]]:
    rows = []
    for r in range(cv.n_rows):
        row = [cell_str(cv.cells.get((r, c))) for c in range(cv.n_cols)]
        if not pad:
            while row and row[-1] == "":
                row.pop()
        rows.append(row)
    return rows


def write_text(cv: Canvas, P: Plan) -> tuple[bytes, str]:
    """Returns (bytes, encoding actually used)."""
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=P.delimiter, quotechar='"', quoting=csv.QUOTE_MINIMAL,
                   lineterminator=P.line_ending)
    w.writerows(canvas_rows(cv, P.pad_rows))
    text = buf.getvalue()
    enc = P.encoding
    if enc == "cp1252":
        try:
            return text.encode("cp1252"), enc
        except UnicodeEncodeError:
            enc = "utf-8"
    if enc == "utf-8-sig":
        return b"\xef\xbb\xbf" + text.encode("utf-8"), enc
    if enc == "utf-16":
        return b"\xff\xfe" + text.encode("utf-16-le"), "utf-16-le-bom"
    return text.encode("utf-8"), "utf-8"


def normalize_zip(data: bytes, fixups: dict | None = None) -> bytes:
    fixups = fixups or {}
    zin = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as zout:
        for info in zin.infolist():
            content = zin.read(info.filename)
            if info.filename in fixups:
                content = fixups[info.filename](content)
            zi = zipfile.ZipInfo(info.filename, date_time=FIXED_ZIP_TIME)
            zi.compress_type = info.compress_type
            zi.external_attr = 0o600 << 16
            zi.create_system = 0
            zout.writestr(zi, content)
    return out.getvalue()


def _fix_core(content: bytes) -> bytes:
    s = content.decode("utf-8")
    s = re.sub(r"(<dcterms:(created|modified)[^>]*>)[^<]*(</dcterms:\2>)", lambda m: m.group(1) + FIXED_ISO + m.group(3), s)
    return s.encode("utf-8")


def _fix_meta(content: bytes) -> bytes:
    s = content.decode("utf-8")
    s = re.sub(r"(<meta:creation-date>)[^<]*(</meta:creation-date>)", r"\g<1>2026-10-09T08:00:00\g<2>", s)
    s = re.sub(r"(<dc:date>)[^<]*(</dc:date>)", r"\g<1>2026-10-09T08:00:00\g<2>", s)
    return s.encode("utf-8")


def write_xlsx(cv: Canvas, P: Plan) -> bytes:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = P.sheet_name
    for (r, c), v in sorted(cv.cells.items()):
        cell = ws.cell(row=r + 1, column=c + 1, value=v)
        if isinstance(v, dt.date):
            cell.number_format = P.xlsx_date_numfmt
    for (r1, c1, r2, c2) in cv.merges:
        if (r1, c1) != (r2, c2):
            ws.merge_cells(start_row=r1 + 1, start_column=c1 + 1, end_row=r2 + 1, end_column=c2 + 1)
    for k in range(P.extra_sheets):
        wb.create_sheet(f"Sheet{k + 2}")
    wb.properties.creator = "spotter-gen"
    wb.properties.created = FIXED_DT
    wb.properties.modified = FIXED_DT
    buf = io.BytesIO()
    wb.save(buf)
    return normalize_zip(buf.getvalue(), {"docProps/core.xml": _fix_core})


XLS_FORMATS = {"m/d/yyyy": "M/D/YYYY", "mm-dd-yy": "MM-DD-YY", "d-mmm": "D-MMM", "ddd m/d": "DDD M/D",
               "yyyy-mm-dd": "YYYY-MM-DD", "mmm d": "MMM D"}


def write_xls(cv: Canvas, P: Plan) -> bytes:
    import xlwt

    wb = xlwt.Workbook(encoding="utf-8")
    ws = wb.add_sheet(P.sheet_name[:31])
    dstyle = xlwt.easyxf(num_format_str=XLS_FORMATS.get(P.xlsx_date_numfmt, "M/D/YYYY"))
    merged_tl = {}
    covered = set()
    for (r1, c1, r2, c2) in cv.merges:
        if (r1, c1) == (r2, c2):
            continue
        merged_tl[(r1, c1)] = (r2, c2)
        for r in range(r1, r2 + 1):
            for c in range(c1, c2 + 1):
                covered.add((r, c))
    for (r1, c1), (r2, c2) in sorted(merged_tl.items()):
        v = cv.cells.get((r1, c1), "")
        if isinstance(v, dt.date):
            ws.write_merge(r1, r2, c1, c2, v, dstyle)
        else:
            ws.write_merge(r1, r2, c1, c2, v)
    for (r, c), v in sorted(cv.cells.items()):
        if (r, c) in covered:
            continue
        if isinstance(v, dt.date):
            ws.write(r, c, v, dstyle)
        else:
            ws.write(r, c, v)
    for k in range(P.extra_sheets):
        wb.add_sheet(f"Sheet{k + 2}")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def write_ods(cv: Canvas, P: Plan) -> bytes:
    from odf.opendocument import OpenDocumentSpreadsheet
    from odf.table import CoveredTableCell, Table, TableCell, TableRow
    from odf.text import P as TP

    doc = OpenDocumentSpreadsheet()
    table = Table(name=P.sheet_name)
    spans = {}
    covered = set()
    for (r1, c1, r2, c2) in cv.merges:
        if (r1, c1) == (r2, c2):
            continue
        spans[(r1, c1)] = (r2 - r1 + 1, c2 - c1 + 1)
        for r in range(r1, r2 + 1):
            for c in range(c1, c2 + 1):
                if (r, c) != (r1, c1):
                    covered.add((r, c))
    for r in range(cv.n_rows):
        tr = TableRow()
        for c in range(cv.n_cols):
            if (r, c) in covered:
                tr.addElement(CoveredTableCell())
                continue
            kw = {}
            if (r, c) in spans:
                kw = {"numberrowsspanned": spans[(r, c)][0], "numbercolumnsspanned": spans[(r, c)][1]}
            v = cv.cells.get((r, c))
            if v is None:
                tc = TableCell(**kw)
            elif isinstance(v, (int, float)) and not isinstance(v, bool):
                tc = TableCell(valuetype="float", value=v, **kw)
                tc.addElement(TP(text=str(v)))
            elif isinstance(v, dt.date):
                tc = TableCell(valuetype="date", datevalue=v.isoformat(), **kw)
                tc.addElement(TP(text=cell_str(v)))
            else:
                tc = TableCell(valuetype="string", **kw)
                for line in str(v).split("\n"):
                    tc.addElement(TP(text=line))
            tr.addElement(tc)
        table.addElement(tr)
    doc.spreadsheet.addElement(table)
    for k in range(P.extra_sheets):
        t = Table(name=f"Sheet{k + 2}")
        t.addElement(TableRow())
        doc.spreadsheet.addElement(t)
    buf = io.BytesIO()
    doc.write(buf)
    return normalize_zip(buf.getvalue(), {"meta.xml": _fix_meta})


def write_docx(bands: list[Band], P: Plan) -> bytes:
    import docx

    d = docx.Document()
    for band in bands:
        for _ in range(band.gap_before):
            d.add_paragraph("")
        for blk in band.blocks:
            if blk.role in ("title", "footer", "blank") and blk.width <= 1:
                for row in blk.rows:
                    d.add_paragraph(cell_str(row[0]) if row else "")
                continue
            lm = band.col_gaps[0]
            w = blk.width + lm
            table = d.add_table(rows=blk.height, cols=w)
            table.style = "Table Grid"
            for (r1, c1, r2, c2) in blk.merges:
                if (r1, c1) != (r2, c2):
                    table.cell(r1, c1 + lm).merge(table.cell(r2, c2 + lm))
            for r, trow in enumerate(table.rows):
                cells = trow.cells
                src = blk.rows[r]
                for c, v in enumerate(src):
                    if v is None or v == "":
                        continue
                    cells[c + lm].text = cell_str(v)
    cp = d.core_properties
    cp.author = "spotter-gen"
    cp.last_modified_by = "spotter-gen"
    cp.created = FIXED_DT
    cp.modified = FIXED_DT
    cp.revision = 1
    buf = io.BytesIO()
    d.save(buf)
    return normalize_zip(buf.getvalue(), {"docProps/core.xml": _fix_core})
